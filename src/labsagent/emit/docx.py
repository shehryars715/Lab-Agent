"""`.docx` -- the Word report, in either of two modes.

ANCHORED MODE is the original design and stays untouched: copy the uploaded
manual and insert each task's work beneath its anchor, so the professor's
numbering, headers, logos and styles survive by construction rather than being
regenerated. `annotate_manual` does that and this emitter only calls it, which
is what keeps the output byte-identical to every report produced before the
seam existed.

FRESH MODE exists because the input is no longer required to be a .docx. A PDF,
a notebook or a pasted message produces no paragraph anchors at all, and the
in-place annotator cannot work without them -- it used to raise `IndexError`
after all the solving money had been spent. So when there are no anchors this
builds a plain document from the block IR instead. It is deliberately modest:
the anchored mode inherits a whole document's design for free, and no
from-scratch layout will match that, so this aims to be clean rather than to
compete.

Choosing between them is not a setting. Anchors exist or they do not.
"""

from __future__ import annotations

from pathlib import Path

import docx as pydocx
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt

from labsagent.blocks import has_screenshot_twin
from labsagent.present import arrange
from labsagent.emit import EmitContext, register
from labsagent.ingest.readers import produces_anchors
from labsagent.report.cover import build_cover
from labsagent.report.docx_builder import (
    BODY_PT,
    CODE_FILL,
    CODE_PT,
    IMAGE_WIDTH_IN,
    annotate_manual,
)
from labsagent.report.docx_utils import shade, style_code_run, tighten


def _label(doc, text: str):
    para = doc.add_paragraph()
    para.add_run(text).bold = True
    tighten(para, before=6, after=2)
    return para


def _code(doc, source: str, size_pt: float = CODE_PT) -> None:
    for line in (source.rstrip("\n").split("\n") or [""]):
        para = doc.add_paragraph()
        run = para.add_run(line if line.strip() else " ")
        style_code_run(run, size_pt=size_pt)
        shade(para, CODE_FILL)
        tighten(para)
        para.paragraph_format.left_indent = Inches(0.25)


def _prose(doc, text: str, size_pt: float = BODY_PT) -> None:
    para = doc.add_paragraph()
    para.add_run(text).font.size = Pt(size_pt)
    tighten(para, before=2, after=8)


def build_fresh(ctx: EmitContext, out_path: Path) -> Path:
    """A report for a run that had no source document to annotate."""
    doc = pydocx.Document()
    pres = ctx.presentation
    body_pt, code_pt = pres.body_pt or BODY_PT, pres.code_pt or CODE_PT

    for position, outcome in enumerate(ctx.outcomes):
        task = outcome.task
        words = ctx.words_for(task.id)
        include = ctx.plan_for(task.id).include
        if position and pres.page_breaks:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(task.title or task.id, level=1)
        if include.statement and task.statement.strip():
            _prose(doc, task.statement.strip(), body_pt)

        wrote_output_label = False
        answers_headed = False
        blocks = arrange(outcome, ctx.style_for(task.id), include)
        for index, block in enumerate(blocks):
            if block.kind == "code":
                if block.title:
                    doc.add_heading(block.title, level=3)
                elif words.code:
                    _label(doc, words.code)
                _code(doc, block.text, code_pt)
                wrote_output_label = False
            elif block.kind == "error":
                if words.status:
                    _label(doc, words.status)
                _prose(doc, block.text, body_pt)
            elif block.kind == "image" and block.path is not None:
                if not wrote_output_label and words.output:
                    _label(doc, words.output)
                wrote_output_label = True
                if Path(block.path).exists():
                    para = doc.add_paragraph()
                    para.add_run().add_picture(
                        str(block.path), width=Inches(IMAGE_WIDTH_IN)
                    )
                    tighten(para, before=2, after=4)
            elif block.kind == "output" and not has_screenshot_twin(blocks, index):
                # Only fall back to the raw text when this output has no
                # terminal picture -- decided per output, not per task, so a
                # section layout shows each section's picture. (Classic used
                # to hide the text whenever ANY figure existed, which lost the
                # output of a run whose screenshot was missing.)
                if not wrote_output_label and words.output:
                    _label(doc, words.output)
                wrote_output_label = True
                for line in block.text.splitlines():
                    para = doc.add_paragraph()
                    style_code_run(para.add_run(line or " "), size_pt=code_pt)
                    tighten(para)
            elif block.kind == "prose":
                if block.role == "answer" and words.answers and not answers_headed:
                    doc.add_heading(words.answers, level=2)
                    answers_headed = True
                if block.title:
                    _label(doc, block.title)
                _prose(doc, block.text, body_pt)

    if ctx.cover is not None:
        build_cover(doc, ctx.cover)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


class DocxEmitter:
    name = "docx"
    label = "Report"
    kind = "report"
    extensions = (".docx",)

    def emit(self, ctx: EmitContext) -> list[Path]:
        out_path = ctx.out_dir / f"{ctx.report_name()}.docx"
        # All three conditions matter: a source document, anchors into it,
        # and a source that can actually carry anchors.
        anchored = (
            ctx.manual_path is not None
            and bool(ctx.anchors)
            and produces_anchors(ctx.manual_path)
        )
        if anchored:
            return [
                annotate_manual(
                    ctx.manual_path,
                    out_path,
                    ctx.outcomes,
                    cover=ctx.cover,
                    anchors=ctx.anchors,
                    presentation=ctx.presentation,
                    identity=ctx.identity,
                )
            ]
        return [build_fresh(ctx, out_path)]


register(DocxEmitter())
