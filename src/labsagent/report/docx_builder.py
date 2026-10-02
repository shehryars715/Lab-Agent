"""Annotate a copy of the lab manual in place.

The manual IS the report template. Original text, numbering, headers, logos and
styles are never regenerated, so they survive by construction. That deletes all
document-design work and replaces it with a narrower problem: resolving anchors
and inserting XML correctly.

WHAT GOES IN AND WHERE IS NOT DECIDED HERE. It used to be: every task got
"Code:", its terminal screenshots and an explanation, inserted after the task --
and a task inside a table was deliberately moved out of it, so a manual's own
answer box stayed empty while the work landed underneath. Each task now comes
with an `AnswerPlan` read from the manual and the student's message
(`present.py`): the parts to include, and the manual's boxes, if it has any.
This module only renders the plan:

    a plan with boxes     each part into the box meant for it, with no label the
                          box already shows; parts with no box follow the last box
    a plan without boxes  the parts after the task's anchor, as before
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from labsagent.blocks import LEAD, Block, has_screenshot_twin
from labsagent.models import TaskOutcome
from labsagent.present import AnswerPlan, Vocabulary, arrange, vocabulary
from labsagent.report.cover import CoverInfo, build_cover
from labsagent.report.docx_utils import (
    enclosing_table,
    flatten_paragraphs,
    insert_paragraph_after,
    insert_paragraph_after_table,
    shade,
    style_code_run,
    tighten,
)

CODE_FILL = "F2F2F2"
CODE_LABEL = "Code:"
OUTPUT_LABEL = "Output:"
IMAGE_WIDTH_IN = 6.0


def _label(cursor: Paragraph, text: str) -> Paragraph:
    para = insert_paragraph_after(cursor)
    run = para.add_run(text)
    run.bold = True
    tighten(para, before=6, after=2)
    return para


def _code_block(cursor: Paragraph, source: str) -> Paragraph:
    """One shaded paragraph per line, so the block page-breaks naturally."""
    lines = source.rstrip("\n").split("\n") or [""]
    for line in lines:
        cursor = insert_paragraph_after(cursor)
        # A space keeps empty lines from collapsing the shaded band.
        run = cursor.add_run(line if line.strip() else " ")
        style_code_run(run)
        shade(cursor, CODE_FILL)
        tighten(cursor)
        cursor.paragraph_format.left_indent = Inches(0.25)
    return cursor


def _image(cursor: Paragraph, image_path: Path, width_in: float = IMAGE_WIDTH_IN) -> Paragraph:
    cursor = insert_paragraph_after(cursor)
    cursor.add_run().add_picture(str(image_path), width=Inches(width_in))
    tighten(cursor, before=2, after=4)
    return cursor


def _explanation(cursor: Paragraph, text: str) -> Paragraph:
    cursor = insert_paragraph_after(cursor)
    run = cursor.add_run(text)
    run.font.size = Pt(10)
    tighten(cursor, before=2, after=8)
    return cursor


def _render(
    cursor: Paragraph,
    blocks: list[Block],
    words: Vocabulary,
    *,
    labels: bool = True,
    width_in: float = IMAGE_WIDTH_IN,
) -> Paragraph:
    """Insert `blocks` after `cursor`, in order; return the last paragraph.

    The cursor advances to each new paragraph -- inserting repeatedly after the
    same one reverses the order. `labels=False` is for a box whose own label
    already says what it holds: a "Code" cell does not need "Code:" again.
    """
    labelled_output = False
    answers_headed = False
    for index, block in enumerate(blocks):
        if block.kind == "code":
            if labels or block.title:
                cursor = _label(cursor, block.title or words.code)
            cursor = _code_block(cursor, block.text)
            labelled_output = False
        elif block.kind == "error":
            if labels:
                cursor = _label(cursor, "Status:")
            cursor = _explanation(cursor, block.text)
        elif block.kind == "image" and block.path is not None and Path(block.path).exists():
            if labels and not labelled_output:
                cursor = _label(cursor, words.output)
            labelled_output = True
            cursor = _image(cursor, Path(block.path), width_in)
        elif (
            block.kind == "output"
            and block.text.strip()
            # TERMINAL OUTPUT IS SHOWN AS THE TERMINAL -- when screenshots are
            # included at all. The text appears exactly when no picture of it
            # follows, decided per output, never by a layout choice.
            and not has_screenshot_twin(blocks, index)
        ):
            if labels and not labelled_output:
                cursor = _label(cursor, words.output)
            labelled_output = True
            cursor = _code_block(cursor, block.text)
        elif block.kind == "prose":
            if block.role == "answer" and words.answers and not answers_headed:
                cursor = _label(cursor, words.answers)
                answers_headed = True
            if block.title:
                cursor = _label(cursor, block.title)
            cursor = _explanation(cursor, block.text)
    return cursor


# --- the manual's own boxes --------------------------------------------------

#: Text a manual leaves for the student to replace: "Write your code here".
_PLACEHOLDER = re.compile(
    r"\b(here|write|paste|insert|type|attach|your (?:code|answer|output|solution))\b", re.I
)


def _replaceable(paragraph: Paragraph) -> bool:
    """An empty box, or a placeholder written to be overwritten. A label such
    as "Code:" is neither: the answer goes under it and the label stays."""
    text = paragraph.text.strip()
    return not text or (len(text) <= 120 and bool(_PLACEHOLDER.search(text)))


def _part_of(block: Block) -> str:
    """Which box a block belongs in."""
    if block.kind == "code" or (block.kind == "prose" and block.role == LEAD):
        return "code"
    if block.kind == "prose" and block.role in ("answer", ""):
        return "answer"
    # Output, pictures, an error, and the notes about what was not done.
    return "output"


def _box_width(paragraph: Paragraph) -> float:
    """How wide a picture may be inside this paragraph's cell, in inches.

    A 6-inch screenshot in a 3-inch cell stretches the table off the page.
    """
    node = paragraph._p.getparent()
    while node is not None and node.tag != qn("w:tc"):
        node = node.getparent()
    if node is None:
        return IMAGE_WIDTH_IN
    width = node.find(f"{qn('w:tcPr')}/{qn('w:tcW')}")
    try:
        twips = int(width.get(qn("w:w")))
        unit = width.get(qn("w:type"))
    except (AttributeError, TypeError, ValueError):
        return 5.0
    if unit in (None, "dxa") and twips > 0:
        return max(1.5, min(IMAGE_WIDTH_IN, twips / 1440 - 0.2))
    return 5.0


def _after_box(paragraph: Paragraph, doc) -> Paragraph:
    """Where parts with no box of their own go: right after the last box."""
    table = enclosing_table(paragraph)
    return insert_paragraph_after_table(table, doc) if table is not None else paragraph


def _fill_boxes(doc, boxes: list[tuple], blocks: list[Block], words: Vocabulary) -> None:
    """Each part into the box meant for it; what has no box follows the last.

    A box for a part ("code", "output", "answer") takes that part only, with no
    label of ours. An "any" box takes everything no specific box claimed, in
    order, with labels -- a single "Solution" box needs to say which is which.
    """
    groups: dict[str, list[Block]] = {"code": [], "output": [], "answer": []}
    for block in blocks:
        groups[_part_of(block)].append(block)

    plan: dict[int, tuple[list[Block], bool]] = {}
    claimed: set[str] = set()
    for slot, _ in boxes:
        if slot.part in groups and slot.part not in claimed:
            plan[slot.idx] = (groups[slot.part], False)
            claimed.add(slot.part)
    rest = [b for b in blocks if _part_of(b) not in claimed]
    for slot, _ in boxes:
        if slot.part == "any" and rest:
            plan[slot.idx] = (rest, True)
            rest = []

    last = boxes[-1][1]
    for slot, paragraph in boxes:
        group, labels = plan.get(slot.idx, ([], False))
        if not group:
            continue
        end = _render(paragraph, group, words, labels=labels, width_in=_box_width(paragraph))
        if end is not paragraph and _replaceable(paragraph):
            # The box's own first line was empty or "write your code here":
            # the answer replaces it rather than sitting under a blank line.
            paragraph._p.getparent().remove(paragraph._p)
            if paragraph is last:
                last = end
        elif paragraph is last:
            last = end

    if rest:
        _render(_after_box(last, doc), rest, words)


def annotate_manual(
    manual_path: Path,
    out_path: Path,
    outcomes: list[TaskOutcome],
    cover: CoverInfo | None = None,
    anchors: dict[str, int] | None = None,
    style: str = "classic",
    plans: dict[str, AnswerPlan] | None = None,
) -> Path:
    """Copy the manual and put each task's answer where its plan says.

    A `cover` is prepended AFTER anchors and boxes have been resolved to
    Paragraph objects. Order matters: prepending shifts every integer index in
    the document, but object references survive it -- the same reason anchors
    are resolved before any mutation at all.
    """
    manual_path, out_path = Path(manual_path), Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(manual_path, out_path)

    doc = docx.Document(str(out_path))
    paragraphs = flatten_paragraphs(doc)
    words = vocabulary(style)

    # Resolve every anchor and box to a Paragraph OBJECT before mutating
    # anything. Integer indices shift the moment the first insert lands; object
    # references stay valid across sibling inserts.
    resolved: list[tuple[Paragraph | None, list[tuple], TaskOutcome, AnswerPlan]] = []
    for outcome in outcomes:
        plan = (plans or {}).get(outcome.task.id) or AnswerPlan()
        boxes = [
            (slot, paragraphs[slot.idx]) for slot in plan.slots if 0 <= slot.idx < len(paragraphs)
        ]
        if boxes:
            resolved.append((None, boxes, outcome, plan))
            continue

        idx = (anchors or {}).get(outcome.task.id, -1)
        if not (0 <= idx < len(paragraphs)):
            raise IndexError(
                f"{outcome.task.id}: anchor {idx} outside document "
                f"(0..{len(paragraphs) - 1})"
            )
        anchor = paragraphs[idx]

        # NO BOX: a task wrapped in a table would otherwise have its work
        # inserted INSIDE the cell that holds the task text -- the statement's
        # cell, not an answer box. Escape to body level.
        tbl = enclosing_table(anchor)
        if tbl is not None:
            anchor = insert_paragraph_after_table(tbl, doc)

        resolved.append((anchor, [], outcome, plan))

    if cover is not None:
        build_cover(doc, cover)

    for anchor, boxes, outcome, plan in resolved:
        blocks = arrange(outcome, style, plan.include)
        if boxes:
            _fill_boxes(doc, boxes, blocks, words)
        else:
            _render(anchor, blocks, words)

    doc.save(str(out_path))
    return out_path
