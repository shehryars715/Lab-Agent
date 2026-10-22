"""The manual's own title page: is there one, and where do the student's details go?

WHY. Many university manuals open with a title page of their own -- the
university, the department, the course, the lab, and blank lines for the
student: "Name: ________", "Registration No: ________". The annotator used to
prepend a cover of ours above it, so the submission opened with two title
pages, the second one still blank. That reads as a tool, and it is exactly the
kind of thing a student then has to fix by hand.

So: when the manual has a title page, our cover is off unless asked for
(`present.CoverSpec.show` is None, "auto"), and the student's details go into
the manual's own blanks instead.

WHAT COUNTS AS A TITLE PAGE. Deliberately strict, because a false positive
costs the student their cover: the first page must END in a break, be SPARSE
(a title page is a few lines, not the start of the lab), mention no task, and
carry either two front-matter facts (course, instructor, date ...) or a blank
for the student's name or ID.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from labsagent.ingest.cover import extract_cover_facts
from labsagent.report.docx_utils import flatten_paragraphs

#: How far into the document a title page may end.
MAX_PARAGRAPHS = 40
#: A title page is a few short lines. Past this it is the lab itself.
MAX_TEXT_LINES = 25
MAX_CHARS = 1200

_TASK = re.compile(r"\b(task|exercise|question|problem)\s*(no\.?\s*)?\d+\b", re.I)

#: A line left for the student: the label, an optional colon, then nothing but
#: a blank to write on.
_BLANK_LINE = re.compile(
    r"^\s*(?P<label>"
    r"(?:student(?:'s)?\s+)?name"
    r"|reg(?:istration)?\.?\s*(?:no\.?|number|#)"
    r"|cms(?:\s*id)?"
    r"|roll\s*(?:no\.?|number|#)"
    r"|enrol?ment\s*(?:no\.?|number|#)"
    r"|student\s*id"
    r"|section|class"
    r")\s*[:：\-–]?\s*(?P<blank>[_.…\s]*)$",
    re.I,
)

_FIELD = {
    "name": "name",
    "section": "section",
    "class": "section",
}


def _field_for(label: str) -> str:
    low = label.lower()
    for key, name in _FIELD.items():
        if low.endswith(key):
            return name
    return "cms_id"


@dataclass(frozen=True)
class TitlePage:
    found: bool = False
    #: The student's details the page has a blank for: "name", "cms_id", ...
    blanks: tuple[str, ...] = ()


def _page_break(paragraph: Paragraph) -> str | None:
    """"before" when a new page starts at this paragraph, "after" when one
    starts after it, None when the paragraph sits inside one page.

    A break run before any text (`<w:br w:type="page"/>` then the words, or
    Word's `lastRenderedPageBreak` hint at the paragraph's start) puts the
    whole paragraph on the next page; one after the text ends the page here.
    The rendered hint is only a hint, which is why the page before it also has
    to be sparse and task-free.
    """
    p = paragraph._p
    if p.find(f"{qn('w:pPr')}/{qn('w:pageBreakBefore')}") is not None:
        return "before"
    seen_text = False
    for node in p.iter(qn("w:t"), qn("w:br"), qn("w:lastRenderedPageBreak")):
        if node.tag == qn("w:t"):
            seen_text = seen_text or bool((node.text or "").strip())
            continue
        if node.tag == qn("w:br") and node.get(qn("w:type")) != "page":
            continue
        return "after" if seen_text else "before"
    if p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None:
        return "after"
    return None


def title_paragraphs(doc) -> list[Paragraph] | None:
    """The paragraphs of the manual's title page, or None when it has none."""
    paragraphs = flatten_paragraphs(doc)[:MAX_PARAGRAPHS]
    page: list[Paragraph] = []
    for paragraph in paragraphs:
        where = _page_break(paragraph)
        # A break before anything was written is not the end of a page.
        if where == "before" and any(p.text.strip() for p in page):
            break
        page.append(paragraph)
        if where == "after":
            break
    else:
        return None

    lines = [p.text.strip() for p in page if p.text.strip()]
    if not lines or len(lines) > MAX_TEXT_LINES or sum(map(len, lines)) > MAX_CHARS:
        return None
    if any(_TASK.search(line) for line in lines):
        return None
    facts = extract_cover_facts(lines)
    known = sum(
        1
        for name in ("course", "instructor", "lab_engineer", "date", "department", "section")
        if getattr(facts, name, None)
    )
    if known < 2 and not _blanks(page):
        return None
    return page


def _blanks(page: list[Paragraph]) -> list[tuple[Paragraph, str]]:
    out = []
    for paragraph in page:
        match = _BLANK_LINE.match(paragraph.text)
        if match:
            out.append((paragraph, _field_for(match.group("label"))))
    return out


def read_title_page(path: Path | str | None) -> TitlePage:
    """For an uploaded .docx: does it have a title page, and which blanks?"""
    if path is None or Path(path).suffix.lower() != ".docx":
        return TitlePage()
    try:
        import docx

        page = title_paragraphs(docx.Document(str(path)))
    except Exception:  # noqa: BLE001 -- unreadable means "no title page", never a crash
        return TitlePage()
    if page is None:
        return TitlePage()
    return TitlePage(found=True, blanks=tuple(dict.fromkeys(f for _, f in _blanks(page))))


def _next_empty_cell(paragraph: Paragraph) -> Paragraph | None:
    """In a table row, the empty cell right after this paragraph's cell."""
    tc = paragraph._p.getparent()
    if tc is None or tc.tag != qn("w:tc"):
        return None
    cell = tc.getnext()
    while cell is not None and cell.tag != qn("w:tc"):
        cell = cell.getnext()
    if cell is None:
        return None
    inner = cell.findall(qn("w:p"))
    if not inner or "".join(t.text or "" for t in cell.iter(qn("w:t"))).strip():
        return None
    return Paragraph(inner[0], paragraph._parent)


def fill_identity(doc, identity: dict[str, str]) -> list[str]:
    """Write the student's details into the title page's blanks.

    Returns the fields written. Only blanks are touched: a line that already
    carries a value is the manual's, not ours.
    """
    page = title_paragraphs(doc)
    if not page:
        return []
    written: list[str] = []
    for paragraph, name in _blanks(page):
        value = (identity.get(name) or "").strip()
        if not value or name in written:
            continue
        target = _next_empty_cell(paragraph)
        if target is not None:
            target.add_run(value)
            written.append(name)
            continue
        # Strip the blank ("_____", "....") off the end of the line, keep the
        # label's own formatting, and write the value after it.
        for run in reversed(paragraph.runs):
            stripped = run.text.rstrip("_.… \t")
            run.text = stripped
            if stripped:
                break
        if ":" not in paragraph.text and "：" not in paragraph.text:
            paragraph.add_run(":")
        paragraph.add_run(f" {value}")
        written.append(name)
    return written
