"""Build every input the matrix sweep needs.

Generated rather than committed, for the same reason `tests/fixtures/make_manual.py`
is: `.gitignore` excludes `*.docx`, and a binary nobody can diff is a poor thing
to keep in git when it is fully reproducible from forty lines of code.

The set is chosen to cover the axes that actually broke things:

    lab05.md / .txt / .pdf   the same lab in three containers, so a failure
                             points at the READER and not at the lab
    ds311.ipynb              a notebook handed out AS the assignment, whose
                             task 2 depends on task 1 and whose text says
                             "submit only the .ipynb"
    cv.docx / invoice.md     documents that are emphatically not labs
    photo.xyz                a binary the upload gate must refuse

    uv run python web/fixtures/make_inputs.py [dir]
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

LAB_LINES = [
    "Lab 05 - Strings and Lists",
    "Course: CS-102 Programming Fundamentals",
    "",
    "Task 1: Reverse a string",
    "Write a program that reads a single word from the user and prints it reversed.",
    "Sample run: input hello, output olleh",
    "",
    "Task 2: Count the vowels",
    "Read a word and print how many vowels it contains.",
    "Sample run: input education, output 5",
    "",
    "Task 3: Longest word",
    "Read a sentence and print its longest word. Explain your approach briefly.",
    "Sample run: input the quick brown fox, output quick",
]


def _markdown(path: Path) -> None:
    body = [("# " + ln) if ln.startswith("Lab 05") else ln for ln in LAB_LINES]
    path.write_text("\n".join(body), encoding="utf-8")


def _text(path: Path) -> None:
    path.write_text("\n".join(LAB_LINES), encoding="utf-8")


def _pdf(path: Path) -> None:
    """A real PDF, laid out as text.

    Worth knowing: extraction puts a space inside words it kerns, so the model
    reads "T ask 1" rather than "Task 1". That is not a flaw in the fixture --
    it is what real PDF extraction does, and the sweep is more honest for it.
    """
    import matplotlib

    matplotlib.use("Agg")
    matplotlib.rcParams["pdf.fonttype"] = 42
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    with PdfPages(path) as pdf:
        fig = plt.figure(figsize=(8.27, 11.69))
        y = 0.95
        for line in LAB_LINES:
            fig.text(0.08, y, line or " ", fontsize=11, family="DejaVu Sans", va="top")
            y -= 0.028
        pdf.savefig(fig)
        plt.close(fig)


def _notebook(path: Path) -> None:
    """The DS311 shape: instructions in markdown cells, one empty code cell.

    Two properties are deliberate. Task 2 says "using the list from task 1", so
    scoping to task 2 alone must drag task 1 in. And the header says "Submit
    only the .ipynb file on LMS", so the artifact choice should follow the
    document rather than the default.
    """
    cells = [
        ("markdown", "# DS311 Lab 4 - Descriptive Statistics\n\n"
                     "Course: DS311 Data Mining. Submit only the .ipynb file on LMS.\n"),
        ("markdown", "## Task 1: Build the dataset\n\n"
                     "Create a list of the 10 integers from 1 to 10 and print it.\n"),
        ("markdown", "## Task 2: Summary statistics\n\n"
                     "Using the list from task 1, compute and print the mean, the median "
                     "and the standard deviation.\n"),
        ("code", "# your code here\n"),
    ]
    nb = {
        "cells": [
            {"cell_type": kind, "metadata": {}, "source": [src],
             **({"execution_count": None, "outputs": []} if kind == "code" else {})}
            for kind, src in cells
        ],
        "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python",
                                    "name": "python3"}},
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    path.write_text(json.dumps(nb, indent=1), encoding="utf-8")


def _cv(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("A Student", 0)
    d.add_paragraph("Software Engineer | Lahore, Pakistan")
    d.add_heading("Experience", 1)
    d.add_paragraph("Engineer, 2024 to present. Agent harnesses and evaluation tooling.")
    d.add_heading("Education", 1)
    d.add_paragraph("BS Computer Science, 2024.")
    d.add_heading("Skills", 1)
    d.add_paragraph("Python, JavaScript, React, FastAPI, Docker.")
    d.save(str(path))


def _invoice(path: Path) -> None:
    path.write_text("\n".join([
        "# Invoice INV-2026-0412",
        "",
        "Bill to: Acme Corp, 14 Mill Road, Lahore",
        "Date: 12 September 2026",
        "",
        "| Item | Qty | Unit | Total |",
        "|---|---|---|---|",
        "| Consulting hours | 24 | 8000 | 192000 |",
        "| Server hosting | 1 | 35000 | 35000 |",
        "",
        "Total due: PKR 227,000. Payable within 30 days.",
    ]), encoding="utf-8")


def _junk(path: Path) -> None:
    """PNG magic bytes under an extension nothing claims to read."""
    path.write_bytes(bytes([0x89, 0x50, 0x4E, 0x47]) + b"\x00" * 400)


BUILDERS = {
    "lab05.md": _markdown,
    "lab05.txt": _text,
    "lab05.pdf": _pdf,
    "ds311.ipynb": _notebook,
    "cv.docx": _cv,
    "invoice.md": _invoice,
    "photo.xyz": _junk,
}


def build(directory: Path, force: bool = False) -> Path:
    """Create any missing input. Returns the directory."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    for name, builder in BUILDERS.items():
        target = directory / name
        if force or not target.exists():
            builder(target)
    return directory


def lab_manual() -> Path:
    """The .docx lab, reusing the fixture the unit tests already generate."""
    path = Path("tests/fixtures/lab03_manual.docx")
    if not path.exists():
        sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "tests" / "fixtures"))
        from make_manual import build as build_manual

        path.parent.mkdir(parents=True, exist_ok=True)
        build_manual(path)
    return path


if __name__ == "__main__":
    where = build(Path(sys.argv[1] if len(sys.argv) > 1 else "web/fixtures/inputs"))
    print(f"inputs in {where}:")
    for p in sorted(where.iterdir()):
        print(f"  {p.name:<16} {p.stat().st_size:>8,} bytes")
