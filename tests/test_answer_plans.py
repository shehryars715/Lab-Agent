"""What an answer contains and where it goes: read from the manual, not templated.

The report used to carry the same parts in the same place for every lab --
"Code:", terminal screenshots, an explanation, all inserted after the task, and
a task inside a table was deliberately moved out of it, so a manual's own
answer box stayed empty. These pin the replacement (2026-10-02): code and
output as text by default, screenshots and explanations only when asked for,
and the manual's boxes filled in place.
"""

from __future__ import annotations

from pathlib import Path

import docx

from labsagent.blocks import SCREENSHOT, Block
from labsagent.ingest.docx_reader import read_manual
from labsagent.ingest.labspec import ExtractedSlot, ExtractedTask, _parts, _slots_for
from labsagent.intent import detect_parts
from labsagent.models import Task, TaskOutcome
from labsagent.present import AnswerPlan, Include, Slot, select
from labsagent.report.docx_builder import annotate_manual
from labsagent.report.docx_utils import flatten_paragraphs


def _outcome(task_id="task1", **task):
    return TaskOutcome(
        task=Task(id=task_id, title="Sum", statement="Add two numbers.", **task),
        status="passed",
        code_text="print(5 + 3)",
        blocks=[
            Block("code", text="print(5 + 3)"),
            Block("output", text="Sum = 8"),
            Block("prose", text="I add the two numbers."),
            Block("prose", text="Because 5 + 3 is 8.", role="answer", title="Why 8?"),
        ],
    )


# --- what is included ---------------------------------------------------------


def test_by_default_there_is_no_screenshot_and_no_explanation(tmp_path: Path):
    shot = tmp_path / "s.png"
    blocks = [
        Block("code", text="x"),
        Block("output", text="8"),
        Block("image", path=shot, role=SCREENSHOT),
        Block("prose", text="I add them."),
        Block("prose", text="Not done: the plot.", role="gap"),
        Block("prose", text="Because.", role="answer", title="Why?"),
    ]
    kept = [(b.kind, b.role) for b in select(blocks, Include())]
    assert kept == [("code", ""), ("output", ""), ("prose", "gap"), ("prose", "answer")]


def test_a_theory_answer_is_never_dropped_as_an_explanation():
    outcome = TaskOutcome(
        task=Task(id="task2", title="Compilers", statement="What is a compiler?", needs_code=False),
        status="passed",
        explanation="It translates source code before it runs.",
    )
    from labsagent.present import arrange

    assert [b.text for b in arrange(outcome, "classic", Include())] == [
        "It translates source code before it runs."
    ]


def test_the_students_words_win_and_hide_wins_a_contradiction():
    manual_said = Include(screenshots=True)
    assert manual_said.changed(hide=["screenshots"]).screenshots is False
    assert Include().changed(show=["explanation"]).explanation is True
    assert Include().changed(show=["screenshots"], hide=["screenshots"]).screenshots is False


def test_the_request_backstop_reads_only_the_obvious():
    assert detect_parts("give me the word report with screenshots") == (["screenshots"], [])
    assert detect_parts("no screenshots please, and add explanations") == (
        ["explanation"], ["screenshots"],
    )
    # A question about the subject is not a request for explanations.
    assert detect_parts("explain how the loop works") == ([], [])
    assert _parts(["Screenshot", "plots", "nonsense"]) == ["screenshots", "figures"]


# --- the manual's boxes -------------------------------------------------------


def _boxed_manual(path: Path) -> Path:
    """Task 1 with a two-row answer table: Code | (empty), Output | (empty)."""
    document = docx.Document()
    document.add_paragraph("Task 1: Add two numbers and print the sum.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Code"
    table.cell(1, 0).text = "Output"
    document.add_paragraph("Task 2: Print hello.")
    document.add_paragraph("Write your code here")
    document.save(str(path))
    return path


def test_the_model_can_see_an_empty_answer_cell(tmp_path: Path):
    view = read_manual(_boxed_manual(tmp_path / "m.docx")).as_numbered_text()
    assert "[2] [cell] (empty)" in view and "[1] [cell] Code" in view


def test_boxes_are_checked_before_they_are_believed(tmp_path: Path):
    manual = read_manual(_boxed_manual(tmp_path / "m.docx"))
    # 0 task1 text, 1 Code, 2 (empty), 3 Output, 4 (empty), 5 task2 text, 6 placeholder
    anchors = {"task1": 0, "task2": 5}
    item = ExtractedTask(
        task_number=1, title="Sum", statement="s", anchor_idx=0, anchor_quote="Task 1",
        answer_slots=[
            ExtractedSlot(part="code", idx=2, quote=""),           # an empty cell: kept
            ExtractedSlot(part="output", idx=4, quote=""),         # kept
            ExtractedSlot(part="any", idx=5, quote="Task 2"),      # next task's text: refused
            ExtractedSlot(part="code", idx=1, quote="Wrong quote"),  # quote mismatch: refused
        ],
    )
    claimed: set[int] = set()
    slots = _slots_for(item, 0, ["task1", "task2"], anchors, manual, claimed)
    assert slots == (Slot(idx=2, part="code"), Slot(idx=4, part="output"))

    # The next task's placeholder is its own box; a cell already taken is not.
    second = ExtractedTask(
        task_number=2, title="Hi", statement="s", anchor_idx=5, anchor_quote="Task 2",
        answer_slots=[
            ExtractedSlot(part="code", idx=6, quote="Write your code here"),
            ExtractedSlot(part="code", idx=2, quote=""),
        ],
    )
    assert _slots_for(second, 1, ["task1", "task2"], anchors, manual, claimed) == (
        Slot(idx=6, part="code"),
    )


def test_a_box_named_by_its_label_moves_to_the_empty_cell(tmp_path: Path):
    """Seen live: the model pointed at the "Code" label cell, not the empty
    cell beside it, and the code would have gone in the label column."""
    manual = read_manual(_boxed_manual(tmp_path / "m.docx"))
    item = ExtractedTask(
        task_number=1, title="Sum", statement="s", anchor_idx=0, anchor_quote="Task 1",
        answer_slots=[
            ExtractedSlot(part="code", idx=1, quote="Code"),
            ExtractedSlot(part="output", idx=3, quote="Output"),
        ],
    )
    slots = _slots_for(item, 0, ["task1", "task2"], {"task1": 0, "task2": 5}, manual, set())
    assert slots == (Slot(idx=2, part="code"), Slot(idx=4, part="output"))


def test_each_part_lands_in_its_own_box_and_extras_follow_the_table(tmp_path: Path):
    manual = _boxed_manual(tmp_path / "m.docx")
    plans = {"task1": AnswerPlan(slots=(Slot(2, "code"), Slot(4, "output")))}
    out = annotate_manual(manual, tmp_path / "r.docx", [_outcome()], anchors={"task1": 0}, plans=plans)

    document = docx.Document(str(out))
    table = document.tables[0]
    assert "print(5 + 3)" in table.cell(0, 1).text
    assert "Sum = 8" in table.cell(1, 1).text
    texts = [p.text for p in flatten_paragraphs(document)]
    assert "Code:" not in texts and "Output:" not in texts, "the boxes already say it"
    # The written answer has no box: it goes right after the table, before Task 2.
    after_table = [p.text for p in document.paragraphs]
    assert after_table.index("Because 5 + 3 is 8.") < after_table.index("Task 2: Print hello.")
    # And the explanation nobody asked for is nowhere.
    assert "I add the two numbers." not in texts


def test_a_placeholder_is_replaced_not_answered_underneath(tmp_path: Path):
    manual = _boxed_manual(tmp_path / "m.docx")
    plans = {"task2": AnswerPlan(slots=(Slot(6, "code"),))}
    out = annotate_manual(
        manual, tmp_path / "r.docx", [_outcome("task2")], anchors={"task2": 5}, plans=plans
    )
    texts = [p.text for p in docx.Document(str(out)).paragraphs]
    assert "Write your code here" not in texts and "print(5 + 3)" in texts


def test_without_boxes_the_answer_still_follows_the_task(tmp_path: Path):
    manual = _boxed_manual(tmp_path / "m.docx")
    out = annotate_manual(manual, tmp_path / "r.docx", [_outcome("task2")], anchors={"task2": 5})
    texts = [p.text for p in docx.Document(str(out)).paragraphs]
    start = texts.index("Task 2: Print hello.")
    assert texts[start + 1] == "Code:" and "Sum = 8" in texts[start:]


# --- the writer is paid for only the words that will be shown -----------------


def test_the_writer_runs_only_when_its_words_are_wanted():
    from web.server.pipeline import _only_when_wanted

    calls = []

    def writer(task, code, transcript):
        calls.append(task.id)
        return "words"

    gate = _only_when_wanted(writer, {"task3": AnswerPlan(include=Include(explanation=True))})
    assert gate(Task(id="task1", title="a", statement="s"), "", None) is None
    gate(Task(id="task2", title="b", statement="s", written_questions=["Why?"]), "", None)
    gate(Task(id="task3", title="c", statement="s"), "", None)
    gate(Task(id="task4", title="d", statement="s", needs_code=False), "", None)
    assert calls == ["task2", "task3", "task4"]


def test_add_screenshots_is_a_change_not_a_vague_reply():
    from web.server.revise import FollowUpPlan, to_followup

    followup = to_followup(
        FollowUpPlan(kind="change", show=["screenshots", "bogus"]), ["task1"], "add screenshots"
    )
    assert followup.kind == "change" and followup.show == ["screenshots"]
    assert not followup.model_work, "a rebuild, never a re-solve"
