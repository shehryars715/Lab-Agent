"""The four cover designs.

Kept separate from `test_cover.py`, which pins the `classic` layout's exact
output. These tests are about the *contract every layout must honour* -- the
things that make a cover a cover regardless of how it looks -- plus the one
thing that makes having four of them worth it: that they actually differ.
"""

from __future__ import annotations

import docx
import pytest

from labsagent.ingest.cover import CoverFacts
from labsagent.models import LabSpec
from labsagent.profile import StudentProfile
from labsagent.report.cover import LAYOUTS, CoverInfo, build_cover, cover_from

FULL = CoverInfo(
    lab_number="10",
    lab_title="Neural Networks-Part A",
    course="CS245 Machine Learning",
    student_name="Shehryar",
    cms_id="22F-1234",
    section="BSDS-02A",
    instructor="Dr. Nazia Perwaiz",
    lab_engineer="Alishba Zulfiqar",
    date="8th April 2026",
    department="Faculty of Computing",
    tagline="Backpropagation from scratch, no framework.",
)

MANUAL_LINES = ["Original Lab Manual", "Task 1: do the thing", "Some closing text"]


@pytest.fixture
def manual(tmp_path):
    doc = docx.Document()
    doc.add_heading(MANUAL_LINES[0], level=1)
    for line in MANUAL_LINES[1:]:
        doc.add_paragraph(line)
    path = tmp_path / "manual.docx"
    doc.save(str(path))
    return path


def render(manual, info: CoverInfo):
    doc = docx.Document(str(manual))
    build_cover(doc, info)
    return doc


def all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


def has_page_break(doc) -> bool:
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for p in doc.paragraphs:
        for run in p.runs:
            if run._element.findall(f"{ns}br"):
                return True
    return False


@pytest.mark.parametrize("layout", LAYOUTS)
class TestEveryLayout:
    """The invariants. A new layout that breaks one of these is not a cover."""

    def test_cover_lands_before_the_manual(self, manual, layout):
        doc = render(manual, CoverInfo(**{**FULL.__dict__, "layout": layout}))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert texts.index(MANUAL_LINES[0]) > 0, "the manual must not be first"

    def test_original_content_survives_verbatim(self, manual, layout):
        doc = render(manual, CoverInfo(**{**FULL.__dict__, "layout": layout}))
        after = all_text(doc)
        for line in MANUAL_LINES:
            assert line in after, f"{layout!r} lost {line!r}"

    def test_ends_with_a_page_break(self, manual, layout):
        doc = render(manual, CoverInfo(**{**FULL.__dict__, "layout": layout}))
        assert has_page_break(doc), f"{layout!r} would run into the manual"

    def test_carries_the_student_identity(self, manual, layout):
        doc = render(manual, CoverInfo(**{**FULL.__dict__, "layout": layout}))
        text = all_text(doc)
        assert "Shehryar" in text
        assert "22F-1234" in text

    def test_unknown_fields_are_omitted_not_blank(self, manual, layout):
        """A row reading 'Instructor:' with nothing after it looks broken."""
        sparse = CoverInfo(student_name="S", cms_id="1", layout=layout)
        doc = render(manual, sparse)
        text = all_text(doc)
        assert "Instructor" not in text
        assert "Lab Engineer" not in text

    def test_survives_a_manual_with_no_known_fields(self, manual, layout):
        doc = render(manual, CoverInfo(layout=layout))
        assert "Lab Report" in all_text(doc)

    def test_is_not_empty(self, manual, layout):
        doc = render(manual, CoverInfo(**{**FULL.__dict__, "layout": layout}))
        assert len([p for p in doc.paragraphs if p.text.strip()]) > 3


class TestLayoutsDiffer:
    """Four layouts that render identically are one layout and three aliases."""

    def test_each_layout_produces_a_distinct_document(self, manual):
        rendered = {}
        for layout in LAYOUTS:
            doc = render(manual, CoverInfo(**{**FULL.__dict__, "layout": layout}))
            rendered[layout] = all_text(doc)
        assert len(set(rendered.values())) == len(LAYOUTS), (
            "two layouts produced identical text; they are not distinct designs"
        )

    def test_layouts_use_different_structural_devices(self, manual):
        """Not just different words -- different shapes on the page."""
        classic = render(manual, CoverInfo(**{**FULL.__dict__, "layout": "classic"}))
        rule = render(manual, CoverInfo(**{**FULL.__dict__, "layout": "rule"}))
        banner = render(manual, CoverInfo(**{**FULL.__dict__, "layout": "banner"}))

        assert classic.tables, "classic is a table"
        # `rule` typesets its fields in an unboxed table and leads with the
        # lab title rather than the composite heading.
        assert "Lab 10: Neural Networks-Part A" not in all_text(rule)
        assert "Neural Networks-Part A" in all_text(rule)
        # `banner` tints a paragraph, which classic never does.
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        assert any(
            p._p.find(f".//{ns}shd") is not None for p in banner.paragraphs
        ), "banner should carry a shaded band"


class TestLayoutSelection:
    def test_unknown_layout_falls_back_to_classic(self, manual):
        """A typo in a layout name must not cost you the report."""
        typo = render(manual, CoverInfo(**{**FULL.__dict__, "layout": "neon-disco"}))
        classic = render(manual, CoverInfo(**{**FULL.__dict__, "layout": "classic"}))
        assert all_text(typo) == all_text(classic)

    def test_default_is_classic(self, manual):
        """The CLI and the existing tests depend on this staying true."""
        default = render(manual, FULL)
        explicit = render(manual, CoverInfo(**{**FULL.__dict__, "layout": "classic"}))
        assert all_text(default) == all_text(explicit)

    def test_cover_from_defaults_preserve_previous_behaviour(self):
        spec = LabSpec(lab_number="13", title="K-Means", course=None, tasks=[])
        info = cover_from(spec, StudentProfile(name="S", cms_id="1"), CoverFacts())
        assert info.layout == "classic"
        assert info.tagline == ""

    def test_cover_from_accepts_a_layout_and_tagline(self):
        spec = LabSpec(lab_number="13", title="K-Means", course=None, tasks=[])
        info = cover_from(
            spec, StudentProfile(name="S", cms_id="1"), CoverFacts(),
            layout="split", tagline="Clustering by hand.",
        )
        assert info.layout == "split"
        assert info.tagline == "Clustering by hand."

    def test_blank_layout_falls_back_rather_than_blanking_the_cover(self):
        spec = LabSpec(lab_number="1", title="T", tasks=[])
        info = cover_from(spec, StudentProfile(name="S", cms_id="1"), CoverFacts(), layout="")
        assert info.layout == "classic"


class TestTagline:
    def test_tagline_appears_where_the_layout_supports_it(self, manual):
        doc = render(manual, CoverInfo(**{**FULL.__dict__, "layout": "rule"}))
        assert FULL.tagline in all_text(doc)

    def test_no_tagline_renders_nothing_extra(self, manual):
        with_tag = render(manual, CoverInfo(**{**FULL.__dict__, "layout": "rule"}))
        without = render(manual, CoverInfo(**{**{**FULL.__dict__, "tagline": ""}, "layout": "rule"}))
        assert FULL.tagline in all_text(with_tag)
        assert FULL.tagline not in all_text(without)
