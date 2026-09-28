"""Harness pass 3 (2026-09-26): terminal output shown as the terminal, a
capability gate and a prerequisite gate before any code is written, and a
report writer that cannot print raw JSON. Offline -- no model, no network."""

from __future__ import annotations

import sys
import zipfile
from dataclasses import replace
from pathlib import Path

import docx
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from labsagent import capabilities  # noqa: E402
from labsagent.agent.explainer import build_brief, overview_of, parse_writeup  # noqa: E402
from labsagent.blocks import SCREENSHOT, blocks_for, has_screenshot_twin, Block  # noqa: E402
from labsagent.capabilities import Requirement  # noqa: E402
from labsagent.config import Settings  # noqa: E402
from labsagent.data import Dataset, provenance_block  # noqa: E402
from labsagent.emit import EmitContext  # noqa: E402
from labsagent.emit.archive import ArchiveEmitter  # noqa: E402
from labsagent.emit.docx import build_fresh  # noqa: E402
from labsagent.ingest.docx_reader import read_manual  # noqa: E402
from labsagent.ingest.labspec import extract_labspec  # noqa: E402
from labsagent.intent import Intent, referenced_task_ids, scope  # noqa: E402
from labsagent.models import LabSpec, Task, TaskOutcome, Transcript  # noqa: E402
from labsagent.orchestrator import _section_shots, build_task_prompt  # noqa: E402
from labsagent.prerequisites import (  # noqa: E402
    OMITTED,
    PROVIDED,
    RECREATED,
    Needed,
    Prerequisite,
    attach,
    report_notes,
    solver_lines,
    writer_lines,
)
from labsagent.present import arrange  # noqa: E402
from labsagent.profile import StudentProfile  # noqa: E402
from labsagent.report.docx_builder import annotate_manual  # noqa: E402
from labsagent.runstore import RunStore, _task_from_dict, _task_to_dict  # noqa: E402
from web.server import pipeline  # noqa: E402
from web.server.briefing import Plan  # noqa: E402
from web.server.jobs import Job  # noqa: E402
from web.server.pipeline import (  # noqa: E402
    answered_notes,
    closing_needs,
    dataset_refs_in,
    prerequisite_questions,
    resolve_prerequisites,
)

LAB02 = Needed(
    what="Your Lab 02 results",
    detail="query customer, filtering decisions",
    task_ids=("task1", "task2"),
    recreatable=True,
    recreate_from="the Online Retail data",
)


def _task(n: int, statement: str = "Do it.", **kw) -> Task:
    return Task(id=f"task{n}", title=f"Task {n}", statement=statement, **kw)


def _png(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (40, 20), "white").save(path)
    return path


class _Shots:
    """A screenshot backend that records what it was asked to draw."""

    def __init__(self) -> None:
        self.calls: list[Transcript] = []

    def render(self, transcript, out_path: Path) -> list[Path]:
        self.calls.append(transcript)
        return [_png(Path(out_path))]


class _Structured:
    """A model whose one structured call returns a fixed extraction."""

    def __init__(self, payload: dict) -> None:
        self.payload = payload

    def with_structured_output(self, schema, method=None):
        payload = self.payload

        class _Bound:
            def invoke(self, prompt, config=None):
                return schema.model_validate(payload)

        return _Bound()


def _manual(tmp_path: Path, *paragraphs: str) -> Path:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    path = tmp_path / "manual.docx"
    document.save(str(path))
    return path


def _extracted_task(n: int, **extra) -> dict:
    return {
        "task_number": n,
        "title": f"Task {n}",
        "statement": f"Statement {n}.",
        "anchor_idx": 0,
        "anchor_quote": "",
        **extra,
    }


# --- cross-document references ------------------------------------------------


def test_another_labs_task_is_not_this_labs_task():
    cases = {
        "A comparison with the nearest customers from Lab 02 Task 4/5.": set(),
        "Compare with Lab 02 Task 4 and Task 5, then extend your Task 1 matrix.": {"task1"},
        "Reuse Task 4 of Lab 02.": set(),
        "Reuse Tasks 4 and 5 from the previous lab.": set(),
        "Use your Task 1 representation and your Lab 02 query customer.": {"task1"},
        "Extend your Task 2 program.": {"task2"},
    }
    for text, want in cases.items():
        assert referenced_task_ids(_task(9, text)) == want, text

    # And so scoping to task 2 no longer drags in this lab's task 4.
    spec = LabSpec("03", "L", tasks=[
        _task(1), _task(2, "Compare with Lab 02 Task 4/5."), _task(3), _task(4),
    ])
    narrowed, pulled = scope(spec, Intent(task_ids=["task2"]))
    assert [t.id for t in narrowed.tasks] == ["task2"] and pulled == []


# --- the capability gate ------------------------------------------------------


def test_the_gate_refuses_what_this_environment_cannot_run():
    spec = LabSpec("02", "L", tasks=[
        _task(1), _task(2), _task(3), _task(4), _task(5),
        _task(6, needs_code=False), _task(7, needs_code=False),
    ])
    requirements = {
        "task1": Requirement(language="html/css"),
        "task2": Requirement(needs=(("display", "a web page to view"),)),
        "task3": Requirement(libraries=("tensorflow",)),
        "task4": Requirement(libraries=("scikit-learn", "json")),  # promised + stdlib
        # Words only: no program, so no language, library or display to lack...
        "task6": Requirement(language="java", needs=(("display", "a GUI"),)),
        # ...but a task done in another TOOL is still out of reach.
        "task7": Requirement(language="none", needs=(("software", "Tableau Prep"),)),
    }
    out = {o.task_id: o for o in capabilities.check(spec, requirements)}
    assert set(out) == {"task1", "task2", "task7"}
    assert "HTML and CSS" in out["task1"].reasons[0]
    assert "screen or browser" in out["task2"].reasons[0]


def test_a_missing_library_is_said_up_front_but_does_not_refuse_the_lab():
    spec = LabSpec("02", "L", tasks=[_task(1), _task(2), _task(3, needs_code=False)])
    requirements = {
        "task1": Requirement(libraries=("colorspacious",)),
        "task2": Requirement(libraries=("sklearn",)),
        "task3": Requirement(libraries=("tensorflow",)),  # words only: nothing to import
    }
    assert capabilities.check(spec, requirements) == []
    found = capabilities.missing_libraries(spec, requirements)
    assert found == [("task1", ["colorspacious"])]
    notice = capabilities.library_notice(found)
    assert notice.startswith("Task 1 asks for the colorspacious library, which isn't installed")
    assert "everything else" in notice
    assert capabilities.library_notice([]) == ""


def test_sklearn_and_the_standard_library_are_available():
    assert capabilities.library_available("sklearn")
    assert capabilities.library_available("matplotlib.pyplot")
    assert capabilities.library_available("tkinter")  # a NEED (display), not a library
    assert not capabilities.library_available("tensorflow")


def test_the_refusal_names_the_task_the_environment_and_how_to_narrow():
    item = capabilities.OutOfScope("task1", "Portfolio site", ["it is written in HTML and CSS"])
    whole = capabilities.refusal([item], total=1)
    assert "Task 1 (Portfolio site)" in whole and "Python 3" in whole
    assert "before writing any code" in whole
    assert "say which ones" not in whole  # nothing to narrow to
    partial = capabilities.refusal([item], total=4)
    assert "say which ones" in partial and "task 1" in partial


def test_every_refusal_reason_must_be_quoted_from_the_manual(tmp_path):
    manual = read_manual(_manual(
        tmp_path,
        "Create a personal portfolio website from scratch using HTML and CSS.",
        "Plot the monthly sales chart and download the dataset from Kaggle.",
        "Train the model with statsmodels.",
        "Write a Python program that generates an HTML report.",
    ))
    tasks = [
        _extracted_task(1, language="html/css", language_evidence="from scratch using HTML and CSS"),
        # Not in the document: a hallucinated language is dropped.
        _extracted_task(2, language="java", language_evidence="Write the solution in Java"),
        _extracted_task(3, needs=[
            {"need": "display", "why": "a chart", "evidence": "Plot the monthly sales chart"},
            {"need": "internet", "why": "data", "evidence": "download the dataset from Kaggle"},
            "server",  # a bare word carries no evidence
        ]),
        _extracted_task(4, libraries=["statsmodels", "tensorflow"]),
        # The model's own quote says Python: the label contradicts it.
        _extracted_task(5, language="html", language_evidence="a Python program that generates an HTML report"),
    ]
    reading = extract_labspec(manual, _Structured({"tasks": tasks}))
    assert reading.requirements["task1"].language == "html/css"
    assert "task2" not in reading.requirements
    assert "task3" not in reading.requirements  # both needs vetoed
    assert reading.requirements["task4"].libraries == ("statsmodels",)
    assert "task5" not in reading.requirements


def test_prerequisites_are_grounded_scoped_to_real_tasks_and_capped(tmp_path):
    manual = read_manual(_manual(
        tmp_path,
        "Briefly restate your Lab 02 filtering decisions and query customer.",
        "Compare with the nearest customers from Lab 02 Task 4/5.",
    ))
    payload = {
        "tasks": [_extracted_task(1), _extracted_task(2)],
        "intent": {"prerequisites": [
            {"what": "Your Lab 02 results", "tasks": [1, "task 2", 9],
             "quote": "restate your Lab 02 filtering decisions", "recreatable": "yes",
             "recreate_from": "the same dataset"},
            {"what": "Invented", "tasks": [1], "quote": "text that is nowhere"},
            {"what": "No tasks", "tasks": [7], "quote": "nearest customers from Lab 02"},
            "a bare string",
        ]},
    }
    needed = extract_labspec(manual, _Structured(payload)).intent.prerequisites
    assert len(needed) == 1
    assert needed[0].task_ids == ("task1", "task2")
    assert needed[0].recreatable and needed[0].recreate_from == "the same dataset"


def test_a_task_naming_the_same_outside_lab_is_attached_even_if_the_model_missed_it(tmp_path):
    manual = read_manual(_manual(tmp_path, "Briefly restate your Lab 02 filtering decisions."))
    payload = {
        "lab_number": "03",
        "tasks": [
            _extracted_task(1),
            _extracted_task(2, statement="Complete this table using results from both Lab 02 and Lab 03."),
            _extracted_task(3, statement="Compare with your Lab 3 Task 1 matrix."),  # its own lab
        ],
        "intent": {"prerequisites": [
            {"what": "Your Lab 02 results", "tasks": [1],
             "quote": "restate your Lab 02 filtering decisions"},
        ]},
    }
    (needed,) = extract_labspec(manual, _Structured(payload)).intent.prerequisites
    assert needed.task_ids == ("task1", "task2")


# --- the prerequisite card ----------------------------------------------------


def test_the_card_offers_only_honest_choices_for_tasks_in_scope():
    spec = LabSpec("03", "L", tasks=[_task(1), _task(2)])
    fixed = replace(LAB02, recreatable=False, recreate_from="")
    kept, fields = prerequisite_questions([LAB02, fixed], spec)
    assert [f["key"] for f in fields] == ["prereq_1", "prereq_2"]
    assert all(f["required"] and f["kind"] == "prerequisite" for f in fields)
    assert [o["value"] for o in fields[0]["options"]] == ["provide", "omit", "recreate"]
    assert [o["value"] for o in fields[1]["options"]] == ["provide", "omit"]
    assert "Tasks 1 and 2 rely on this" in fields[0]["reason"]

    only_task3 = LabSpec("03", "L", tasks=[_task(3)])
    assert prerequisite_questions([LAB02], only_task3) == ([], [])


def test_nothing_short_of_a_complete_answer_lets_the_run_continue():
    spec = LabSpec("03", "L", tasks=[_task(1), _task(2)])
    kept, fields = prerequisite_questions([LAB02], spec)
    assert resolve_prerequisites(kept, fields, {}) is None  # timeout / Stop here
    assert resolve_prerequisites(kept, fields, {"prereq_1": "skip"}) is None
    assert resolve_prerequisites(kept, fields, {"prereq_1": "provide"}) is None
    assert resolve_prerequisites(kept, fields, {"prereq_1": "omit"}) == {0: (OMITTED, "")}
    given = {"prereq_1": "provide", "prereq_1_value": "customer 14646"}
    assert resolve_prerequisites(kept, fields, given) == {0: (PROVIDED, "customer 14646")}

    fixed = [replace(LAB02, recreatable=False)]
    _, no_recreate = prerequisite_questions(fixed, spec)
    assert resolve_prerequisites(fixed, no_recreate, {"prereq_1": "recreate"}) is None


def test_card_answers_never_become_notes_or_downloads():
    answers = {
        "prereq_1": "provide",
        "prereq_1_value": "my notebook: https://example.com/lab02.csv",
        "constants": "g = 9.8",
    }
    notes = answered_notes([{"key": "constants", "label": "Constants"}], answers)
    assert "prereq" not in notes and "g = 9.8" in notes
    assert dataset_refs_in(answers) == []


class _AnsweringJob(Job):
    def __init__(self, answers: dict) -> None:
        super().__init__("test")
        self.canned = answers
        self.asked: list = []

    def ask(self, questions, *, timeout_s):
        self.asked.append(questions)
        return dict(self.canned)

    def events(self) -> list[dict]:
        return self._log


def _drive(monkeypatch, tmp_path, reading, answers=None) -> _AnsweringJob:
    def must_not_solve(*a, **k):  # noqa: ARG001
        raise AssertionError("the run reached the solver")

    monkeypatch.setattr(pipeline, "extract_labspec", lambda *a, **k: reading)
    monkeypatch.setattr(pipeline, "build_model", lambda *a, **k: None)
    monkeypatch.setattr(
        pipeline, "read_briefing",
        lambda *a, **k: (Plan(questions=[], layout="classic", tagline=""), reading.usage),
    )
    monkeypatch.setattr(pipeline, "run_lab", must_not_solve)
    monkeypatch.setattr(pipeline, "RUNS_ROOT", tmp_path / "runs")
    job = _AnsweringJob(answers or {})
    manual = _manual(tmp_path, "Lab 03", "Task 1: do it.")
    pipeline.run_job(
        job, manual_path=manual, instructions="", profile_seed={},
        settings=Settings(deepseek_api_key="x"),
    )
    return job


def _reading(spec, **kw):
    from labsagent.ingest.labspec import Reading
    from labsagent.usage import Usage

    return Reading(kind="lab", confidence=0.9, what_this_is="a lab", usage=Usage(), spec=spec, **kw)


def test_an_html_lab_is_refused_before_briefing_or_any_run_directory(monkeypatch, tmp_path):
    spec = LabSpec("02", "Portfolio", tasks=[_task(1)])
    reading = _reading(spec, intent=Intent(), requirements={"task1": Requirement(language="html/css")})
    job = _drive(monkeypatch, tmp_path, reading)
    assert job.status == "failed" and "I can't do this lab" in job.error
    assert not job.asked
    assert not (tmp_path / "runs").exists()
    assert any(e["type"] == "out_of_scope" for e in job.events())


def test_an_unanswered_prerequisite_stops_the_run_before_any_code(monkeypatch, tmp_path):
    spec = LabSpec("03", "Similarity", tasks=[_task(1), _task(2)])
    reading = _reading(spec, intent=Intent(prerequisites=[LAB02]))
    job = _drive(monkeypatch, tmp_path, reading, answers={})
    assert job.status == "failed" and "relies on work from outside" in job.error
    assert job.asked and job.asked[0][0]["key"] == "prereq_1"
    narration = " ".join(e.get("text", "") for e in job.events() if e["type"] == "narration")
    assert "Tasks 1 and 2 rely on your Lab 02 results, which you didn't send" in narration


def test_an_answered_prerequisite_reaches_the_solved_tasks(monkeypatch, tmp_path):
    spec = LabSpec("03", "Similarity", tasks=[_task(1), _task(2), _task(3)])
    reading = _reading(spec, intent=Intent(prerequisites=[LAB02]))
    seen = {}

    def fake_run_lab(spec, store, *a, **k):
        seen["spec"] = spec
        raise RuntimeError("stop here")  # run_job never raises; it reports it

    job = _drive(monkeypatch, tmp_path, reading, answers={"prereq_1": "recreate"})
    assert "reached the solver" in job.error  # an answer lets the run through

    monkeypatch.setattr(pipeline, "run_lab", fake_run_lab)
    job = _AnsweringJob({"prereq_1": "recreate"})
    pipeline.run_job(
        job, manual_path=_manual(tmp_path, "x"), instructions="", profile_seed={},
        settings=Settings(deepseek_api_key="x"),
    )
    by_id = {t.id: t for t in seen["spec"].tasks}
    assert by_id["task1"].prerequisites[0].resolution == RECREATED
    assert by_id["task1"].prerequisites[0].source == "the Online Retail data"
    assert by_id["task3"].prerequisites == []


# --- one fact, three renderings --------------------------------------------------


def _with(resolution: str, value: str = "") -> Task:
    spec = LabSpec("03", "L", tasks=[_task(1, instruction="\n\nStudent notes win.")])
    decisions = {0: (resolution, value)}
    return attach(spec, [replace(LAB02, task_ids=("task1",))], decisions).tasks[0]


def test_the_solver_is_told_what_it_may_do_before_the_students_own_words():
    plain = _task(1)
    assert solver_lines(plain) == ""
    assert build_task_prompt(plain, {}) == "Do it.\n\nWrite your solution to a file named task1.py"

    omitted = _with(OMITTED)
    prompt = build_task_prompt(omitted, {})
    assert "Do not reconstruct, assume or recompute it" in prompt
    assert prompt.index("Do not reconstruct") < prompt.index("Student notes win.")
    assert prompt.rstrip().endswith("task1.py")
    assert "explicitly asked you to recreate it" in solver_lines(_with(RECREATED))
    assert "customer 14646" in solver_lines(_with(PROVIDED, "customer 14646"))


def test_the_writer_is_told_what_it_may_claim_and_nothing_about_retries():
    for task in (_with(OMITTED), _with(RECREATED), _with(PROVIDED, "customer 14646")):
        brief = build_brief(task, "print(1)", Transcript(command="python task1.py", lines=["1"]))
        assert writer_lines(task).strip() in brief
        lowered = brief.lower()
        for word in ("attempt", "retry", "failed", "fix"):
            assert word not in lowered, word
    assert "not provided" in writer_lines(_with(OMITTED))
    assert "NOT the student's original" in writer_lines(_with(RECREATED))


def test_the_report_labels_a_recreation_in_every_writer(tmp_path):
    task = _with(RECREATED)
    outcome = TaskOutcome(task=task, status="passed", code_text="print(1)")
    notes = report_notes(outcome)
    assert notes and notes[0].startswith("Recreated for this report")
    assert any(b.role == "outside" for b in blocks_for(outcome))

    manual = _manual(tmp_path, "Task 1: restate your Lab 02 results.")
    out = annotate_manual(manual, tmp_path / "r.docx", [outcome], anchors={"task1": 0})
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert text.count("Recreated for this report") == 1


def test_a_theory_task_first_is_not_printed_twice_with_the_data_note(tmp_path):
    theory = TaskOutcome(
        task=_task(1, needs_code=False), status="passed", explanation="The only answer.",
    )
    csv = tmp_path / "d.csv"
    csv.write_text("a\n1\n", encoding="utf-8")
    note = provenance_block([Dataset(name="d.csv", path=csv, origin="upload")])
    theory.blocks = [note, *blocks_for(theory)]
    manual = _manual(tmp_path, "Task 1: discuss.")
    out = annotate_manual(manual, tmp_path / "r.docx", [theory], anchors={"task1": 0})
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert text.count("The only answer.") == 1 and text.count("Data used") == 1


def test_resolved_prerequisites_survive_the_manifest():
    task = _with(PROVIDED, "customer 14646")
    back = _task_from_dict(_task_to_dict(task))
    assert back.prerequisites == task.prerequisites
    old = _task_to_dict(_task(1))
    del old["prerequisites"]
    assert _task_from_dict(old).prerequisites == []
    odd = {**old, "prerequisites": [{"what": "X", "resolution": "weird"}, {"nope": 1}]}
    assert _task_from_dict(odd).prerequisites == [Prerequisite(what="X", resolution=OMITTED)]


def test_a_declined_gap_is_not_asked_for_again():
    declined = TaskOutcome(task=_with(OMITTED), status="passed", gap="Left out the comparison.")
    other = TaskOutcome(task=_task(2), status="passed", gap="No browser for the check.")
    message = closing_needs([declined, other])
    assert "No browser" in message and "Left out the comparison" not in message


# --- the report writer never prints JSON ------------------------------------------


def test_a_write_up_that_will_not_parse_is_salvaged_or_dropped_never_printed():
    qs = ["Q1?", "Q2?"]
    assert parse_writeup("Just prose back.", qs).overview == "Just prose back."
    trailing = parse_writeup('{"overview": "O.", "answers": ["A1.", "A2."]} Hope this helps!', qs)
    assert trailing.overview == "O." and len(trailing.answers) == 2
    newline = parse_writeup('{"overview": "One.\nTwo.", "answers": ["A1.", "A2."]}', qs)
    assert newline.overview == "One. Two."
    quoted = parse_writeup('{"overview": "The "best" one.", "answers": ["A1.", "A2."]}', qs)
    assert quoted.overview == 'The "best" one.'
    cut = parse_writeup('{"overview": "O.", "answers": ["A1.", "A2 is cut', qs)
    assert [a["answer"] for a in cut.answers] == ["A1."]
    assert parse_writeup('{"foo": ', qs) is None
    for reply in ('{"overview": "O.", "answers": [', '{"foo": "bar"}'):
        written = parse_writeup(reply, qs)
        assert written is None or not written.overview.lstrip().startswith("{")
    assert overview_of('{"overview": "Only this."}') == "Only this."


# --- terminal output is shown as the terminal ---------------------------------------


def _sectioned(tmp_path: Path, with_shots: bool) -> TaskOutcome:
    sections = [
        {"title": "setup", "code": "import pandas as pd", "output": "", "figures": []},
        {"title": "load", "code": "print('rows:', n)", "output": "rows: 5", "figures": []},
        {"title": "query", "code": "print('customer:', c)", "output": "customer: 7", "figures": []},
    ]
    if with_shots:
        _section_shots(sections, _task(1), "python task1.py", _Shots(), _Store(tmp_path))
    return TaskOutcome(
        task=_task(1), status="passed", code_text="x", sections=sections,
        transcript=Transcript(command="python task1.py", lines=["rows: 5", "customer: 7"]),
        screenshot_paths=[_png(tmp_path / "whole.png")],
    )


class _Store:
    def __init__(self, root: Path) -> None:
        self.shots_dir = root / "shots"


def test_each_section_gets_a_slice_of_the_one_terminal_session(tmp_path):
    shots = _Shots()
    sections = [
        {"output": ""},
        {"output": "rows: 5"},
        {"output": "customer: 7"},
    ]
    _section_shots(sections, _task(1), "python task1.py", shots, _Store(tmp_path))
    assert "screenshots" not in sections[0]  # printed nothing, pictures nothing
    assert [(t.head, t.tail) for t in shots.calls] == [(True, False), (False, True)]
    assert shots.calls[0].display_lines() == [r"PS C:\lab> python task1.py", "rows: 5"]
    assert sections[2]["screenshots"][0].endswith("task1_s03_output.png")


def test_a_picture_that_cannot_be_drawn_leaves_the_text(tmp_path):
    class Broken:
        def render(self, transcript, out_path):
            raise OSError("no font")

    sections = [{"output": "rows: 5"}]
    _section_shots(sections, _task(1), "python task1.py", Broken(), _Store(tmp_path))
    assert "screenshots" not in sections[0]


def _body_text(path: Path) -> str:
    return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)


def _image_count(path: Path) -> int:
    # Placed pictures, not media files: python-docx stores identical PNGs once.
    return len(docx.Document(str(path)).inline_shapes)


def _ctx(tmp_path, outcome, style) -> EmitContext:
    return EmitContext(
        spec=LabSpec("03", "L", tasks=[outcome.task]), outcomes=[outcome],
        out_dir=tmp_path, profile=StudentProfile(), style=style,
    )


def test_a_section_layout_shows_terminal_pictures_not_text(tmp_path):
    outcome = _sectioned(tmp_path, with_shots=True)
    blocks = arrange(outcome, "findings")
    outputs = [i for i, b in enumerate(blocks) if b.kind == "output"]
    assert outputs and all(has_screenshot_twin(blocks, i) for i in outputs)

    fresh = build_fresh(_ctx(tmp_path, outcome, "findings"), tmp_path / "fresh.docx")
    assert "rows: 5" not in _body_text(fresh) and _image_count(fresh) == 2

    manual = _manual(tmp_path, "Task 1: load the data.")
    anchored = annotate_manual(
        manual, tmp_path / "anchored.docx", [outcome], anchors={"task1": 0}, style="findings"
    )
    assert "customer: 7" not in _body_text(anchored) and _image_count(anchored) == 2


def test_sections_from_before_the_pictures_still_show_their_text(tmp_path):
    outcome = _sectioned(tmp_path, with_shots=False)
    fresh = build_fresh(_ctx(tmp_path, outcome, "walkthrough"), tmp_path / "legacy.docx")
    assert "rows: 5" in _body_text(fresh)


def test_a_missing_picture_file_does_not_swallow_the_output(tmp_path):
    blocks = [Block("output", text="x"), Block("image", path=tmp_path / "gone.png", role=SCREENSHOT)]
    assert not has_screenshot_twin(blocks, 0)


def test_the_zip_carries_the_section_pictures(tmp_path):
    outcome = _sectioned(tmp_path, with_shots=True)
    out = tmp_path / "zip"
    out.mkdir()
    ctx = replace(_ctx(out, outcome, "findings"), produced=[])
    (path,) = ArchiveEmitter().emit(ctx)
    names = zipfile.ZipFile(path).namelist()
    assert "screenshots/task1_s02_output.png" in names and "screenshots/whole.png" in names


def test_the_run_store_keeps_section_pictures(tmp_path):
    outcome = _sectioned(tmp_path, with_shots=True)
    store = RunStore.create("03", root=tmp_path / "runs")
    from labsagent.models import RunManifest
    from datetime import datetime, timezone

    manifest = RunManifest(run_id=store.run_id, started_at=datetime.now(timezone.utc),
                           spec=LabSpec("03", "L", tasks=[outcome.task]), outcomes=[outcome])
    store.save(manifest)
    loaded = store.load().outcomes[0]
    assert loaded.sections[1]["screenshots"] == outcome.sections[1]["screenshots"]
