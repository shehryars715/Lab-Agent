"""Tests for the in-place annotation engine.

The fidelity test is the one that mechanically proves "the agent doesn't touch
anything" -- the promise the whole report design rests on.
"""

from __future__ import annotations

from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn

from labsagent.capture.rendered import RenderedBackend
from labsagent.models import Task, TaskOutcome, Transcript
from labsagent.report.docx_builder import annotate_manual
from labsagent.report.docx_utils import flatten_paragraphs

BLIP = ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"


def _signature(path: Path) -> list[tuple[str, str]]:
    return [(p.text, p.style.name) for p in flatten_paragraphs(docx.Document(str(path)))]


def _outcome(tmp_path: Path, task_id: str = "task1") -> TaskOutcome:
    shot = RenderedBackend().render(
        Transcript(command=f"python {task_id}.py", lines=["Enter n: 5", "Sum = 8"]),
        tmp_path / f"{task_id}.png",
    )
    return TaskOutcome(
        task=Task(id=task_id, title="T", statement="s"),
        status="passed",
        code_text='n = int(input("Enter n: "))\nif n:\n    print(f"Sum = {n}")\n',
        screenshot_paths=shot,
        explanation="An explanation.",
        attempts=1,
    )


def test_flatten_includes_table_cell_paragraphs(manual_path):
    """doc.paragraphs silently skips table cells; many manuals wrap tasks in one."""
    doc = docx.Document(str(manual_path))
    flat = flatten_paragraphs(doc)

    assert len(flat) > len(doc.paragraphs)
    assert any("reads 5 integers" in p.text for p in flat)
    assert not any("reads 5 integers" in p.text for p in doc.paragraphs)


def test_every_original_paragraph_survives(manual_path, tmp_path):
    before = _signature(manual_path)
    out = annotate_manual(manual_path, tmp_path / "r.docx", [_outcome(tmp_path)], anchors={"task1": 8})
    after = _signature(out)

    missing = [s for s in before if s not in after]
    assert missing == [], f"annotation lost original content: {missing}"
    assert len(after) > len(before)


def test_inserted_sequence_is_in_order(manual_path, tmp_path):
    """Catches the reverse-order XML bug directly."""
    out = annotate_manual(manual_path, tmp_path / "r.docx", [_outcome(tmp_path)], anchors={"task1": 8})
    flat = flatten_paragraphs(docx.Document(str(out)))

    texts = [p.text for p in flat]
    code_at = texts.index("Code:")
    output_at = texts.index("Output:")

    assert code_at < output_at
    assert texts[code_at + 1].startswith("n = int(input")
    assert texts[code_at + 3].strip().startswith("print(")  # source order preserved
    assert flat[output_at + 1]._p.findall(BLIP), "screenshot missing after Output:"
    assert flat[output_at + 2].text == "An explanation."


def test_insertion_lands_under_the_right_task(manual_path, tmp_path):
    out = annotate_manual(
        manual_path,
        tmp_path / "r.docx",
        [_outcome(tmp_path, "task1"), _outcome(tmp_path, "task2")],
        anchors={"task1": 8, "task2": 10},
    )
    texts = [p.text for p in flatten_paragraphs(docx.Document(str(out)))]

    first_task = texts.index("Task 1: Sum of Two Numbers")
    second_task = texts.index("Task 2: Even Numbers")
    code_blocks = [i for i, t in enumerate(texts) if t == "Code:"]

    assert len(code_blocks) == 2
    assert first_task < code_blocks[0] < second_task < code_blocks[1]


def test_code_indentation_is_preserved(manual_path, tmp_path):
    """Python is whitespace-significant; Word must not collapse it."""
    out = annotate_manual(manual_path, tmp_path / "r.docx", [_outcome(tmp_path)], anchors={"task1": 8})
    flat = flatten_paragraphs(docx.Document(str(out)))

    indented = [p for p in flat if p.text.startswith("    print(")]
    assert indented, "indented code line lost its leading whitespace"

    node = indented[0]._p.findall(".//" + qn("w:t"))[0]
    assert node.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"


def test_code_is_shaded_and_monospaced(manual_path, tmp_path):
    out = annotate_manual(manual_path, tmp_path / "r.docx", [_outcome(tmp_path)], anchors={"task1": 8})
    flat = flatten_paragraphs(docx.Document(str(out)))
    code_line = next(p for p in flat if p.text.startswith("n = int(input"))

    shd = code_line._p.find(qn("w:pPr")).find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == "F2F2F2"
    assert code_line.runs[0].font.name == "Consolas"


def test_failed_task_is_annotated_not_skipped(manual_path, tmp_path):
    failed = TaskOutcome(
        task=Task(id="task1", title="T", statement="s"),
        status="failed",
        code_text="print(undefined_name)\n",
        attempts=3,
        error="NameError: name 'undefined_name' is not defined",
    )
    out = annotate_manual(manual_path, tmp_path / "r.docx", [failed], anchors={"task1": 8})
    texts = [p.text for p in flatten_paragraphs(docx.Document(str(out)))]

    assert "Code:" in texts
    assert any("did not complete after 3 attempt" in t for t in texts)
    assert any("NameError" in t for t in texts)


def test_out_of_range_anchor_is_rejected_loudly(manual_path, tmp_path):
    with pytest.raises(IndexError, match="anchor"):
        annotate_manual(manual_path, tmp_path / "r.docx", [_outcome(tmp_path)], anchors={"task1": 9999})


def test_source_manual_is_never_modified(manual_path, tmp_path):
    before = manual_path.read_bytes()
    annotate_manual(manual_path, tmp_path / "r.docx", [_outcome(tmp_path)], anchors={"task1": 8})

    assert manual_path.read_bytes() == before


def test_table_wrapped_task_inserts_after_the_table(manual_path, tmp_path):
    """Anchor 13 is inside the Task 3 table cell. Inserting there would cram a
    6in screenshot into a bordered box and look unlike every other task."""
    out = annotate_manual(manual_path, tmp_path / "r.docx", [_outcome(tmp_path)], anchors={"task1": 13})
    doc = docx.Document(str(out))

    from docx.table import Table

    from labsagent.report.docx_utils import iter_block_items

    table = next(b for b in iter_block_items(doc) if isinstance(b, Table))
    cell_text = [p.text for p in table.rows[0].cells[0].paragraphs]

    assert "Code:" not in cell_text, "work leaked into the table cell"
    assert any("reads 5 integers" in t for t in cell_text), "original cell text lost"

    body = [b.text for b in iter_block_items(doc) if not isinstance(b, Table)]
    assert "Code:" in body, "work did not land at body level"
