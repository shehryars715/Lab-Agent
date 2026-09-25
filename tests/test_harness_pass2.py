"""Harness pass 2 (2026-09-25): stop instead of improvising, fetch the linked
dataset, ask only when blocked, don't polish. Offline -- no model, no network."""

from __future__ import annotations

import importlib
import sys
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from langchain_core.messages import AIMessage

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fakes import ScriptedModel  # noqa: E402
from labsagent import events as ev  # noqa: E402
from labsagent.agent.build import fast_variant, solver_config  # noqa: E402
from labsagent.agent.prompts import AVAILABLE_LIBRARIES, SOLVER_PROMPT  # noqa: E402
from labsagent.agent.tools import TaskRecorder, build_tools, missing_module  # noqa: E402
from labsagent.budget import LiveUsage  # noqa: E402
from labsagent.config import Settings  # noqa: E402
from labsagent.data.preview import detect_encoding, profile  # noqa: E402
from labsagent.data.refs import harvest_refs, reconcile_datasets  # noqa: E402
from labsagent.data.sources import (  # noqa: E402
    classify,
    kaggle_competition,
    kaggle_slug,
    unreadable_reason,
)
from labsagent.ingest.docx_reader import read_manual  # noqa: E402
from labsagent.intent import Intent  # noqa: E402
from labsagent.models import LabSpec, Task, TaskOutcome  # noqa: E402
from labsagent.orchestrator import error_tail, plain_sentence, run_lab, solve_task  # noqa: E402
from labsagent.runstore import RunStore  # noqa: E402
from labsagent.usage import RunUsage  # noqa: E402
from web.server.briefing import Briefing, ProposedQuestion, _to_plan  # noqa: E402
from web.server.pipeline import closing_needs, data_question  # noqa: E402

KAGGLE = "https://www.kaggle.com/datasets/vivek468/superstore-dataset-final"


class _Shots:
    def render(self, transcript, out_path: Path) -> list[Path]:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_bytes(b"")
        return [out_path]


def _linked_manual(tmp_path: Path) -> Path:
    """The real Lab 2 shape: the words are linked, the URL is nowhere in the text."""
    document = docx.Document()
    paragraph = document.add_paragraph("Download ")
    rid = document.part.relate_to(KAGGLE, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Superstore Dataset"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    paragraph.add_run(" from Kaggle and use it for this lab task.")
    document.add_paragraph("Task A: plot a histogram of Quantity.")
    path = tmp_path / "linked.docx"
    document.save(path)
    return path


# --- 2. Kaggle links in Word files ---------------------------------------------


def test_a_hyperlink_target_reaches_the_model_without_moving_the_text(tmp_path):
    manual = read_manual(_linked_manual(tmp_path))
    first = manual.paragraphs[0]
    assert first.text == "Download Superstore Dataset from Kaggle and use it for this lab task."
    assert first.links == (("Superstore Dataset", KAGGLE),)
    assert f"(link: Superstore Dataset -> {KAGGLE})" in manual.as_numbered_text()


def test_the_harvester_finds_dataset_links_and_ignores_documentation():
    texts = [
        "Read the tutorial at https://colorspacious.readthedocs.io/en/latest/tutorial.html.",
        f"({KAGGLE}).",
        "Or run: kaggle datasets download -d uciml/iris",
        "Grab https://example.edu/data/sales.csv?raw=1 too.",
    ]
    assert harvest_refs(texts) == [
        KAGGLE,
        "uciml/iris",
        "https://example.edu/data/sales.csv?raw=1",
    ]


def test_a_model_slug_the_document_never_mentions_is_dropped():
    texts = ["Download Car Specification Dataset from Kaggle.", KAGGLE.replace(
        "vivek468/superstore-dataset-final", "jahaidulislam/car-specification-dataset-1945-2020")]
    kept = reconcile_datasets(["vivek468/superstore-dataset-final"], texts)
    assert kept == [texts[1]]


def test_kaggle_reference_shapes():
    assert kaggle_slug("www.kaggle.com/datasets/uciml/iris") == "uciml/iris"
    assert kaggle_slug("(https://www.kaggle.com/uciml/iris).") == "uciml/iris"
    assert kaggle_slug("mirichoi0218/insurance.") == "mirichoi0218/insurance"
    assert kaggle_slug("https://www.kaggle.com/code/someone/a-notebook") is None
    assert kaggle_competition("https://www.kaggle.com/c/titanic") == "titanic"
    assert classify("kaggle competitions download -c titanic") == "competition"


def test_a_windows_1252_csv_is_profiled_with_its_encoding(tmp_path):
    path = tmp_path / "superstore.csv"
    path.write_bytes("Region,Sales\nSouth,261.96\nWest\xa0Coast,14.62\n".encode("cp1252"))
    assert detect_encoding(path) == "cp1252"
    text = profile(path)
    assert "columns: Region" in text and 'encoding="cp1252"' in text


def test_formats_nothing_installed_can_read_are_refused_in_plain_words():
    for name in ("old.xls", "table.parquet"):
        reason = unreadable_reason(Path(name))
        engine = "xlrd" if name.endswith(".xls") else "pyarrow"
        if importlib.util.find_spec(engine) is None:
            assert reason and "Please" in reason
    assert unreadable_reason(Path("fine.xlsx")) is None


# --- 1. a sanctioned stop ---------------------------------------------------------


def test_every_library_the_prompt_promises_imports():
    """A prompt that advertises a capability is a contract with the environment."""
    for name in AVAILABLE_LIBRARIES:
        assert name in SOLVER_PROMPT
        importlib.import_module({"scikit-learn": "sklearn"}.get(name, name))


def test_the_recorder_tool_ends_the_loop_and_accepts_blocked(tmp_path):
    class _Box:
        workdir = tmp_path

    recorder = TaskRecorder()
    record = build_tools(_Box(), recorder)[1]
    assert record.return_direct
    record.invoke({"entry_file": "task1.py", "status": "blocked", "missing": "Needs X."})
    assert (recorder.status, recorder.missing) == ("blocked", "Needs X.")


def test_a_missing_module_is_named_unless_the_workspace_supplies_it(tmp_path):
    stderr = "Traceback ...\nModuleNotFoundError: No module named 'notreal_xyz.sub'"
    assert missing_module(stderr, tmp_path) == "notreal_xyz"
    (tmp_path / "notreal_xyz.py").write_text("", encoding="utf-8")
    assert missing_module(stderr, tmp_path) is None
    assert missing_module("ModuleNotFoundError: No module named 'json'", tmp_path) is None


def _blocked_script() -> list[AIMessage]:
    return [
        AIMessage(content="", tool_calls=[{
            "name": "run_solution", "args": {"entry_file": "task1.py", "stdin_values": []},
            "id": "c1"}]),
        AIMessage(content="", tool_calls=[{
            "name": "record_task_result",
            "args": {"entry_file": "task1.py", "status": "blocked",
                     "missing": "This task needs `colorspacious`, which isn't installed here."},
            "id": "c2"}]),
        AIMessage(content="never read: return_direct ended the loop"),
    ]


def test_a_declared_blocker_is_not_retried_and_says_why(tmp_path):
    task = Task(id="task1", title="Colour", statement="Simulate colour blindness.")
    store = RunStore.create("99", root=tmp_path / "runs")
    workspace = store.workspace / task.id
    workspace.mkdir(parents=True, exist_ok=True)
    (workspace / "task1.py").write_text("import notreal_xyz\n", encoding="utf-8")
    model = ScriptedModel(script=_blocked_script())

    outcome = solve_task(
        task, store, Settings(deepseek_api_key=""), _Shots(), RunUsage(), ev.Emitter(),
        {}, model=model,
    ).outcome

    assert outcome.status == "failed" and outcome.attempts == 1
    assert outcome.blocker == "This task needs colorspacious, which isn't installed here."
    assert outcome.runs == 1
    assert model.calls == 2  # no retry, and no extra turn after the tool


def test_a_task_built_on_a_blocked_task_is_not_started(tmp_path, monkeypatch):
    import labsagent.orchestrator as orch

    calls = []

    def fake_solve(task, *args, **kwargs):
        calls.append(task.id)
        outcome = TaskOutcome(task=task, status="failed", blocker="Needs X.")
        return orch.SolveResult(outcome=outcome, cost_usd=0.0)

    monkeypatch.setattr(orch, "solve_task", fake_solve)
    spec = LabSpec(lab_number="1", title="L", tasks=[
        Task(id="task1", title="A", statement="Load it."),
        Task(id="task2", title="B", statement="Extend task 1 with a chart."),
    ])
    store = RunStore.create("1", root=tmp_path / "runs")
    manifest = run_lab(spec, store, Settings(deepseek_api_key=""), _Shots())
    assert calls == ["task1"]
    assert "builds on task1" in manifest.outcomes[1].blocker


def test_the_retry_hint_keeps_the_exception_line():
    trace = "Traceback (most recent call last):\n" + '  File "x.py", line 1\n' * 40 + (
        "KeyError: 'Species'"
    )
    assert error_tail(trace).endswith("KeyError: 'Species'")
    assert plain_sentence("Traceback (most recent call last):\nNeeds `seaborn`.") == "Needs seaborn."


def test_the_task_ceiling_spans_attempts():
    usage = RunUsage()
    live = LiveUsage(usage.phase("solve"), run_usage=usage, task_cap=0.02, task_start=-0.05)
    assert live.task_spent == 0.05


# --- 4. speed for basic tasks ---------------------------------------------------------


def test_a_fast_attempt_copies_the_model_with_thinking_off():
    from langchain_deepseek import ChatDeepSeek

    settings = Settings(deepseek_api_key="")
    model = ChatDeepSeek(model="deepseek-flash", api_key="sk-test")
    fast = fast_variant(model, settings)
    assert fast.reasoning_effort == "none" and model.reasoning_effort is None
    scripted = ScriptedModel(script=[AIMessage(content="")])
    assert fast_variant(scripted, settings) is scripted
    assert solver_config(settings, max_turns=12)["recursion_limit"] == 26


def test_effort_survives_the_manifest(tmp_path):
    store = RunStore.create("1", root=tmp_path / "runs")
    task = Task(id="task1", title="A", statement="s", effort="basic")
    manifest = run_lab(LabSpec(lab_number="1", title="L", tasks=[]), store,
                       Settings(deepseek_api_key=""), _Shots())
    manifest.spec = LabSpec(lab_number="1", title="L", tasks=[task])
    manifest.outcomes = [TaskOutcome(task=task, status="failed", blocker="b", gap=None, runs=3)]
    store.save(manifest)
    loaded = store.load()
    assert loaded.spec.tasks[0].effort == "basic"
    assert (loaded.outcomes[0].blocker, loaded.outcomes[0].runs) == ("b", 3)


# --- 3. ask only on a critical blocker ----------------------------------------------


def test_a_question_that_blocks_no_task_is_dropped():
    briefing = Briefing(questions=[
        ProposedQuestion(key="fmt", label="Plain text or a table?"),
        ProposedQuestion(key="k", label="Which constants?", blocks=["task2"]),
        ProposedQuestion(key="k2", label="And these?", blocks=["task1"]),
    ])
    plan = _to_plan(briefing, set(), ["task1", "task2"])
    assert [q["label"] for q in plan.questions] == ["Which constants?"]


def test_the_data_question_is_asked_only_when_nothing_can_be_fetched():
    unlinked = Intent(data_unlinked=["the Superstore dataset from Kaggle"])
    question = data_question(unlinked, [])
    assert question["key"] == "datasets" and question["required"]
    assert "Superstore" in question["reason"]
    assert data_question(Intent(datasets=[KAGGLE], data_unlinked=["x"]), []) is None
    assert data_question(unlinked, [Path("upload.csv")]) is None
    assert data_question(Intent(), []) is None


def test_the_closing_message_lists_every_stop_and_gap():
    a = TaskOutcome(task=Task(id="task1", title="Task A", statement=""), status="failed",
                    blocker="Needs seaborn.")
    b = TaskOutcome(task=Task(id="task2", title="Task C", statement=""), status="passed",
                    gap="Left out the colour-blindness part.")
    text = closing_needs([a, b])
    assert text.startswith("To finish, I need:")
    assert "- Task A: Needs seaborn." in text and "- Task C: Left out" in text
    assert closing_needs([]) == ""
