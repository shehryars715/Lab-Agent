"""The flexibility pass (2026-10-22): one report spec, receipts, the cover,
the manual's own title page, new formats, changes after a restart, and the
editing agent's guard and replay.

No model is called anywhere here. The editing agent's scripts are written by
hand and run through the real local sandbox, which is the part that has to
be right: what ships is what a saved script produces, checked by code.
"""

from __future__ import annotations

import zipfile
from datetime import datetime, timezone
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_BREAK
from PIL import Image

from labsagent.config import load_settings
from labsagent.emit import EmitContext, emit_all
from labsagent.ingest.cover import CoverFacts
from labsagent.intent import detect_formats, detect_parts
from labsagent.models import LabSpec, RunManifest, Task, TaskOutcome, Transcript
from labsagent.present import (
    AnswerPlan,
    CoverPatch,
    CoverSpec,
    Include,
    LookPatch,
    Presentation,
    Slot,
    TaskLookPatch,
    apply,
    is_empty,
    receipt,
    split_identity,
)
from labsagent.profile import StudentProfile
from labsagent.report.cover import cover_from
from labsagent.report.docx_builder import annotate_manual
from labsagent.report.docx_utils import flatten_paragraphs
from labsagent.report.title_page import fill_identity, read_title_page
from labsagent.runstore import RunStore, manifest_from_dict, manifest_to_dict

CODE = "total = 5 + 3\nprint('Sum =', total)"

# -- helpers -------------------------------------------------------------------


def _settings():
    return load_settings().model_copy(update={"sandbox_backend": "local"})


def _outcome(task_id="task1", title="Sum", statement="Add two numbers and print the sum."):
    return TaskOutcome(
        task=Task(id=task_id, title=title, statement=statement),
        status="passed",
        code_text=CODE,
        transcript=Transcript(command="python task1.py", lines=["Sum = 8"], prompt="PS>"),
        answers=[{"question": "Why 8?", "answer": "Because five and three make eight."}],
    )


def _spec(*outcomes):
    return LabSpec(lab_number="03", title="Loops", tasks=[o.task for o in outcomes])


def _report(path: Path, picture: Path | None = None) -> Path:
    """A finished Word report: a heading, the code, the output, an answer."""
    doc = docx.Document()
    doc.add_heading("Sum", level=1)
    for line in CODE.splitlines():
        doc.add_paragraph(line)
    doc.add_paragraph("Sum = 8")
    doc.add_paragraph("Because five and three make eight.")
    if picture is not None:
        doc.add_picture(str(picture))
    doc.save(str(path))
    return path


def _png(path: Path) -> Path:
    Image.new("RGB", (40, 20), "white").save(path)
    return path


def _texts(path: Path) -> list[str]:
    return [p.text for p in flatten_paragraphs(docx.Document(str(path)))]


# -- 1. the report spec ------------------------------------------------------------


def test_a_cover_patch_sets_overrides_adds_and_removes_lines():
    before = Presentation()
    patch = LookPatch(
        cover=CoverPatch(
            layout="banner",
            title="My Loops Lab",
            fields={"teacher": "Dr. Ali", "submitted to": "Sir Ahmed", "Date": ""},
        )
    )
    after = apply(before, patch, ["task1"])
    assert after.cover.layout == "banner"
    assert after.cover.title == "My Loops Lab"
    # "teacher" is the Instructor line; a name the cover has no line for adds one.
    assert after.cover.fields == {"Instructor": "Dr. Ali", "Submitted to": "Sir Ahmed", "Date": ""}

    said = receipt(before, after, ["task1"])
    assert "Cover design: banner." in said
    assert "The cover now says Instructor: Dr. Ali." in said
    assert "The cover now says Submitted to: Sir Ahmed." in said
    assert "Took Date off the cover." in said


def test_the_cover_info_honours_the_students_lines():
    spec = LabSpec(lab_number="03", title="Loops", tasks=[])
    facts = CoverFacts(course="CS-101", instructor="Prof. Wrong", date="1 Jan")
    info = cover_from(
        spec,
        StudentProfile(name="Ayesha", cms_id="22F-1"),
        facts,
        title="My title",
        fields={"Instructor": "Dr. Ali", "Date": "", "Submitted to": "Sir Ahmed"},
    )
    rows = dict(info.rows())
    assert rows["Instructor"] == "Dr. Ali"
    assert "Date" not in rows
    assert rows["Submitted to"] == "Sir Ahmed"
    assert info.heading() == "My title"
    # The split layout puts a student's own line with the course lines.
    assert ("Submitted to", "Sir Ahmed") in info.course_rows()


def test_who_the_student_is_rides_on_the_profile_not_the_cover():
    patch, identity = split_identity(
        LookPatch(cover=CoverPatch(fields={"Name": "Ayesha Abid", "Reg No": "22F-1", "Instructor": "Dr. Ali"}))
    )
    assert identity == {"name": "Ayesha Abid", "cms_id": "22F-1"}
    assert patch.cover.fields == {"Instructor": "Dr. Ali"}


def test_the_students_words_win_over_the_briefings_taste():
    taste = Presentation(cover=CoverSpec(layout="rule", tagline="Model's line."), style="walkthrough")
    asked = LookPatch(cover=CoverPatch(layout="classic", show=False), style="compact")
    after = apply(taste, asked, ["task1"])
    assert after.cover.layout == "classic" and after.cover.show is False
    assert after.style == "compact"
    assert after.cover.tagline == "Model's line."  # untouched: nobody asked


def test_cover_auto_is_off_only_for_a_manual_with_its_own_title_page():
    assert CoverSpec().shown(manual_has_title_page=False)
    assert not CoverSpec().shown(manual_has_title_page=True)
    assert CoverSpec(show=True).shown(manual_has_title_page=True)
    said = receipt(
        Presentation(), apply(Presentation(), LookPatch(cover=CoverPatch(show="true"))), title_page=True
    )
    assert said == ["Added a cover."]


def test_labels_fonts_breaks_formats_and_names_are_validated():
    after = apply(
        Presentation(),
        LookPatch(
            labels={"Code": "Source code", "output": ""},
            body_pt="40",
            code_pt="7.3pt",
            page_breaks="yes",
            formats=["PDF", ".html", "exe"],
            file_name="Ayesha Abid Lab 3.docx",
        ),
        ["task1"],
    )
    assert after.labels == {"code": "Source code", "output": ""}
    assert after.words_for("task1").code == "Source code"
    assert after.words_for("task1").output == ""
    assert after.body_pt == 16.0 and after.code_pt == 7.5
    assert after.page_breaks is True
    assert after.formats == ("pdf", "html")
    assert after.stem == "Ayesha_Abid_Lab_3"

    reset = apply(after, LookPatch(labels={"code": None}), ["task1"])
    assert "code" not in reset.labels
    assert "The code label is back to normal." in receipt(after, reset, ["task1"])


def test_a_change_to_one_task_touches_only_that_task():
    after = apply(
        Presentation(),
        LookPatch(tasks=[TaskLookPatch(task_ids=["task2", "task9"], hide=["code"], style="compact")]),
        ["task1", "task2"],
    )
    assert after.include_for("task1").code is True
    assert after.include_for("task2").code is False
    assert after.style_for("task2") == "compact" and after.style_for("task1") == "classic"
    assert "Left out the code in task 2." in receipt(Presentation(), after, ["task1", "task2"])


def test_an_empty_or_unknown_patch_is_empty():
    assert is_empty(None) and is_empty(LookPatch())
    assert is_empty(LookPatch(style="nonsense", formats=["exe"]))
    assert not is_empty(LookPatch(cover=CoverPatch(show=False)))


def test_the_spec_survives_the_manifest_and_old_manifests_still_load():
    pres = Presentation(
        cover=CoverSpec(show=False, layout="split", title="T", tagline="L", fields={"Date": ""}),
        style="findings",
        tasks={
            "task1": AnswerPlan(
                include=Include(screenshots=True, statement=False), slots=(Slot(4, "code"),), style="compact"
            )
        },
        labels={"code": "Program"},
        page_breaks=True,
        body_pt=11.0,
        formats=("docx", "pdf"),
        stem="Mine",
        edits=({"n": 1, "request": "blue headings", "target": "docx", "script": "edits/01.py", "status": "applied"},),
    )
    assert Presentation.from_dict(pres.to_dict()) == pres
    assert Presentation.from_dict({}) == Presentation()

    manifest = RunManifest(
        run_id="r", started_at=datetime.now(timezone.utc), spec=_spec(_outcome()),
        presentation=pres.to_dict(), context={"manual": "input/m.docx"},
    )
    back = manifest_from_dict(manifest_to_dict(manifest))
    assert Presentation.from_dict(back.presentation) == pres
    assert back.context == {"manual": "input/m.docx"}
    old = manifest_to_dict(manifest)
    del old["presentation"], old["context"]
    assert manifest_from_dict(old).presentation == {}


# -- 2. the router -----------------------------------------------------------------


def test_a_look_wish_with_no_field_goes_to_the_editor_not_vague():
    from web.server.revise import FollowUpPlan, to_followup

    followup = to_followup(FollowUpPlan(kind="change", edit_request="make the headings dark blue"), ["task1"])
    assert followup.kind == "change"
    assert followup.edit_request == "make the headings dark blue"
    assert not followup.model_work


def test_a_name_typed_as_a_cover_line_reaches_the_profile():
    from web.server.revise import FollowUpPlan, to_followup

    followup = to_followup(
        FollowUpPlan(kind="change", look=LookPatch(cover=CoverPatch(fields={"Name": "Ayesha"}))), ["task1"]
    )
    assert followup.kind == "change"
    assert followup.identity == {"name": "Ayesha"}
    assert followup.look is None


def test_the_fields_decide_not_the_label():
    from web.server.revise import FollowUpPlan, to_followup

    # Seen live: a name, labelled as a question.
    named = to_followup(
        FollowUpPlan(kind="answer", reply="Noted.", identity={"name": "Ayesha", "cms_id": "22F-1"}), ["task1"]
    )
    assert named.kind == "change" and named.identity == {"name": "Ayesha", "cms_id": "22F-1"}
    # What no file can do changes nothing, and is said.
    lms = to_followup(
        FollowUpPlan(kind="change", reply="I'll submit it.", unsupported=["submit it on LMS"]), ["task1"]
    )
    assert lms.is_answer and lms.reply == "I can't submit it on LMS.", "no rebuild for nothing"
    asked = to_followup(FollowUpPlan(kind="answer", unsupported=["submit it on LMS"]), ["task1"])
    assert asked.is_answer and asked.reply == "I can't submit it on LMS."


def test_the_router_reads_how_the_files_look_now():
    from web.server.revise import describe_look

    pres = apply(
        Presentation(),
        LookPatch(cover=CoverPatch(layout="banner", fields={"Instructor": "Dr. Ali"}), labels={"code": ""}),
    )
    text = describe_look(pres, default_formats=("docx", "zip"))
    assert "layout banner" in text and "Instructor: Dr. Ali" in text
    assert "code -> (none)" in text and "docx (the Word report)" in text
    assert "Cover: none (the manual has its own title page)" in describe_look(Presentation(), title_page=True)


# -- 3. the backstop's words -----------------------------------------------------


@pytest.mark.parametrize(
    "text,show,hide",
    [
        ("remove the screenshots", [], ["screenshots"]),
        ("screenshots are not needed", [], ["screenshots"]),
        ("exclude screenshots please", [], ["screenshots"]),
        ("no pictures at all", [], ["screenshots", "figures"]),
        ("drop the charts", [], ["figures"]),
        ("skip the task statements", [], ["statement"]),
        ("plot the sales by region", [], []),
        ("with screenshots and no explanations", ["screenshots"], ["explanation"]),
    ],
)
def test_the_backstop_reads_refusals_in_either_position(text, show, hide):
    assert detect_parts(text) == (show, hide)


def test_the_new_formats_are_recognised_and_refusals_are_not():
    assert set(detect_formats("also a pdf and a web page")) == {"pdf", "html"}
    assert detect_formats("no pdf please") == []
    assert detect_formats("remove the zip") == []


# -- 4. the manual's own title page ---------------------------------------------------


def _manual_with_title_page(path: Path, *, task_on_cover=False) -> Path:
    doc = docx.Document()
    doc.add_paragraph("Department of Computing")
    doc.add_paragraph("Course: CS-101 Programming Fundamentals")
    doc.add_paragraph("Instructor: Dr. Sana")
    if task_on_cover:
        doc.add_paragraph("Task 1: write a program")
    doc.add_paragraph("Name: ____________")
    doc.add_paragraph("Registration No: ____________")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Task 1", level=1)
    doc.add_paragraph("Write a program that adds two numbers.")
    doc.save(str(path))
    return path


def test_a_title_page_is_found_with_its_blanks(tmp_path):
    page = read_title_page(_manual_with_title_page(tmp_path / "m.docx"))
    assert page.found and page.blanks == ("name", "cms_id")


def test_a_first_page_that_holds_a_task_is_not_a_title_page(tmp_path):
    assert not read_title_page(_manual_with_title_page(tmp_path / "m.docx", task_on_cover=True)).found
    plain = docx.Document()
    plain.add_paragraph("Lab 3. Task 1: add two numbers.")
    plain.save(str(tmp_path / "plain.docx"))
    assert not read_title_page(tmp_path / "plain.docx").found
    assert not read_title_page(None).found


def test_with_no_cover_of_ours_the_details_go_into_the_manuals_blanks(tmp_path):
    manual = _manual_with_title_page(tmp_path / "m.docx")
    anchor = next(i for i, t in enumerate(_texts(manual)) if t.startswith("Write a program"))
    out = annotate_manual(
        manual,
        tmp_path / "r.docx",
        [_outcome()],
        cover=None,
        anchors={"task1": anchor},
        identity={"name": "Ayesha Abid", "cms_id": "22F-1"},
    )
    texts = _texts(out)
    assert "Name: Ayesha Abid" in texts
    assert "Registration No: 22F-1" in texts
    assert texts.count("Lab 03: Loops") == 0  # no cover of ours above theirs


def test_fill_identity_never_overwrites_a_line_that_has_a_value(tmp_path):
    doc = docx.Document(str(_manual_with_title_page(tmp_path / "m.docx")))
    assert fill_identity(doc, {"name": "A", "cms_id": "1", "section": "B"}) == ["name", "cms_id"]
    assert fill_identity(doc, {"name": "Someone else"}) == []


# -- 5. what the spec changes in the files -------------------------------------------


def test_labels_page_breaks_and_sizes_reach_the_annotated_report(tmp_path):
    manual = _manual_with_title_page(tmp_path / "m.docx")
    anchor = next(i for i, t in enumerate(_texts(manual)) if t.startswith("Write a program"))
    second = _outcome("task2", title="Twice")
    pres = Presentation(labels={"code": "", "output": "Result"}, page_breaks=True, code_pt=11.0)
    out = annotate_manual(
        manual, tmp_path / "r.docx", [_outcome(), second],
        anchors={"task1": anchor, "task2": anchor}, presentation=pres,
    )
    texts = _texts(out)
    assert "Code:" not in texts and "Result" in texts
    doc = docx.Document(str(out))
    breaks = sum(
        1 for p in doc.paragraphs for r in p.runs
        for br in r._element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br")
        if br.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page"
    )
    assert breaks == 2  # the manual's own, and one between the two tasks
    code_runs = [r for p in doc.paragraphs for r in p.runs if r.text == "total = 5 + 3"]
    assert code_runs and code_runs[0].font.size.pt == 11.0


def _ctx(tmp_path, pres, outcomes=None, cover=True):
    outcomes = outcomes or [_outcome()]
    spec = _spec(*outcomes)
    profile = StudentProfile(name="Ayesha Abid", cms_id="22F-1")
    return EmitContext(
        spec=spec,
        outcomes=outcomes,
        out_dir=tmp_path / "report",
        profile=profile,
        cover=cover_from(spec, profile, CoverFacts(course="CS-101"), fields=pres.cover.fields) if cover else None,
        presentation=pres,
    )


def test_every_new_format_is_built_and_honours_the_spec(tmp_path):
    pres = Presentation(
        labels={"code": "Source code"},
        tasks={"task1": AnswerPlan(include=Include(statement=False))},
        cover=CoverSpec(fields={"Submitted to": "Sir Ahmed"}),
        stem="Ayesha_Lab3",
    )
    ctx = _ctx(tmp_path, pres)
    errors = []
    written = emit_all(ctx, ["md", "pdf", "html", "txt", "files", "zip"], on_error=lambda n, e: errors.append((n, e)))
    assert errors == []
    names = sorted(p.name for p in written)
    assert names == sorted(
        ["Ayesha_Lab3.md", "Ayesha_Lab3.pdf", "Ayesha_Lab3.html", "Ayesha_Lab3.txt",
         "Ayesha_Lab3_task1.py", "Ayesha_Lab3.zip"]
    )
    out = tmp_path / "report"
    md = (out / "Ayesha_Lab3.md").read_text(encoding="utf-8")
    assert "Add two numbers" not in md, "the statement was left out"
    html = (out / "Ayesha_Lab3.html").read_text(encoding="utf-8")
    assert "Source code" in html and "Sir Ahmed" in html
    txt = (out / "Ayesha_Lab3.txt").read_text(encoding="utf-8")
    assert "Source code" in txt and "Sum = 8" in txt

    from pypdf import PdfReader

    pdf_text = " ".join(page.extract_text() for page in PdfReader(str(out / "Ayesha_Lab3.pdf")).pages)
    assert "Sir Ahmed" in pdf_text and "Source code" in pdf_text and "Sum = 8" in pdf_text

    task_file = (out / "tasks" / "Ayesha_Lab3_task1.py").read_text(encoding="utf-8")
    assert task_file.startswith("# Sum") and "Ayesha Abid" in task_file and CODE in task_file

    with zipfile.ZipFile(out / "Ayesha_Lab3.zip") as zf:
        names = zf.namelist()
    # The task files ARE the code folder; no header-less duplicates beside them.
    assert "code/Ayesha_Lab3_task1.py" in names
    assert not any(n.startswith("code/") and n != "code/Ayesha_Lab3_task1.py" for n in names)


def test_page_breaks_put_each_task_on_its_own_pdf_page(tmp_path):
    from pypdf import PdfReader

    two = [_outcome(), _outcome("task2", title="Twice")]
    for breaks, pages in ((False, 1), (True, 2)):
        ctx = _ctx(tmp_path / str(breaks), Presentation(page_breaks=breaks), two, cover=False)
        (path,) = emit_all(ctx, ["pdf"])
        assert len(PdfReader(str(path)).pages) == pages


# -- 6. changes after a restart -----------------------------------------------------


def _finished_run(tmp_path, pres=None):
    """A run folder as `run_job` leaves it: manifest, context, files."""
    from labsagent.capture.rendered import RenderedBackend
    from labsagent.report.title_page import TitlePage
    from web.server.pipeline import RunContext

    store = RunStore.create("03", root=tmp_path / "runs", suffix="job1")
    outcome = _outcome()
    manual = store.dir / "input" / "m.docx"
    manual.parent.mkdir(parents=True)
    _manual_with_title_page(manual)
    # No anchors on the manifest here on purpose: `save_to` must write them.
    manifest = RunManifest(
        run_id=store.run_id, started_at=datetime.now(timezone.utc), spec=_spec(outcome),
        outcomes=[outcome],
    )
    context = RunContext(
        store=store,
        spec=manifest.spec,
        facts=CoverFacts(course="CS-101", instructor="Dr. Sana"),
        profile=StudentProfile(name="Ayesha", cms_id="22F-1"),
        manual_path=manual,
        settings=_settings(),
        presentation=pres or Presentation(style="findings", formats=("docx", "zip")),
        anchors={"task1": 8},
        screenshots=RenderedBackend(look="notebook", folder="Lab 03"),
        title_page=TitlePage(found=True, blanks=("name",)),
    )
    _report(store.report_dir / "Lab03_Report.docx")
    context.save_to(manifest)
    store.save(manifest)
    return store, context


def test_a_run_context_is_rebuilt_from_its_folder_alone(tmp_path):
    from web.server.pipeline import RunContext

    store, saved = _finished_run(tmp_path)
    # The account says otherwise; the files were made for 22F-1, and keep it.
    back = RunContext.from_disk(store, _settings(), StudentProfile(cms_id="99X-0", section="B"))
    assert (back.profile.name, back.profile.cms_id, back.profile.section) == ("Ayesha", "22F-1", "B")
    assert back.presentation == saved.presentation
    assert back.facts.course == "CS-101" and back.facts.instructor == "Dr. Sana"
    assert back.manual_path == store.dir / "input" / "m.docx"
    assert back.anchors == {"task1": 8}
    assert back.screenshots.look == "notebook" and back.screenshots.prompt.endswith("Lab 03>")
    assert back.title_page.found and back.title_page.blanks == ("name",)


def test_a_lab_from_history_is_reopened_as_a_finished_job(tmp_path):
    from web.server.jobs import Job
    from web.server.pipeline import reopen_job

    store, _ = _finished_run(tmp_path)
    job = Job("reopened")
    reopen_job(job, store.dir, _settings(), {"name": "Ayesha"})
    assert job.is_terminal() and job.context is not None
    assert job.context.store.run_id == store.run_id
    assert "docx" in job.artifacts
    assert job.last_seq >= 3  # spec, the file, the finish


# -- 7. the editing agent's guard and replay ----------------------------------------

BLUE_HEADINGS = """import docx
from docx.shared import RGBColor
doc = docx.Document("report.docx")
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
doc.add_paragraph("Submitted on time.")
doc.save("report.docx")
print("headings blue")
"""

DROPS_CODE = """import docx
doc = docx.Document("report.docx")
for p in doc.paragraphs:
    if p.text.startswith("total ="):
        p._p.getparent().remove(p._p)
doc.save("report.docx")
print("tidied")
"""

DROPS_PICTURE = """import docx
doc = docx.Document("report.docx")
for p in doc.paragraphs:
    for d in p._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'):
        d.getparent().remove(d)
doc.save("report.docx")
print("tidied")
"""


def test_the_guard_keeps_a_look_change_and_rejects_lost_work(tmp_path):
    from web.server.edit import check_edit

    before = _report(tmp_path / "before.docx", picture=_png(tmp_path / "p.png"))
    styled = docx.Document(str(before))
    styled.paragraphs[0].runs[0].bold = True
    styled.add_paragraph("A line the student dictated.")
    styled.save(str(tmp_path / "styled.docx"))
    assert check_edit(before, tmp_path / "styled.docx", [_outcome()]) is None

    lost = docx.Document(str(before))
    lost.paragraphs[1]._p.getparent().remove(lost.paragraphs[1]._p)
    lost.save(str(tmp_path / "lost.docx"))
    assert "changed part of your work" in check_edit(before, tmp_path / "lost.docx", [_outcome()])

    (tmp_path / "broken.docx").write_bytes(b"not a document")
    assert "would not open" in check_edit(before, tmp_path / "broken.docx", [_outcome()])


@pytest.mark.parametrize(
    "source,expected",
    [
        (BLUE_HEADINGS, None),
        (DROPS_CODE, "changed part of your work"),
        (DROPS_PICTURE, "removed 1 picture"),
        ("raise SystemExit(3)\n", "stopped with an error"),
    ],
)
def test_a_saved_script_is_run_by_us_and_checked(tmp_path, source, expected):
    from web.server.edit import apply_script

    store = RunStore.create("03", root=tmp_path / "runs", suffix="e")
    target = _report(store.report_dir / "Lab03_Report.docx", picture=_png(tmp_path / "p.png"))
    original = target.read_bytes()
    reason = apply_script(_settings(), store, source, target, [_outcome()])
    if expected is None:
        assert reason is None
        assert "Submitted on time." in _texts(target)
    else:
        assert expected in reason
        assert target.read_bytes() == original, "a rejected edit leaves the file as it was"


def test_saved_edits_are_replayed_after_a_rebuild_and_a_stale_one_is_said(tmp_path):
    from labsagent.usage import RunUsage
    from web.server.jobs import Job
    from web.server.pipeline import _emit

    store, context = _finished_run(tmp_path)
    (store.dir / "edits").mkdir()
    (store.dir / "edits" / "01.py").write_text(BLUE_HEADINGS, encoding="utf-8")
    (store.dir / "edits" / "02.py").write_text(DROPS_CODE, encoding="utf-8")
    context.manual_path = None  # a fresh report: the replay is what is under test
    context.title_page = context.title_page.__class__()
    context.presentation = Presentation(
        formats=("docx", "zip"),
        edits=(
            {"n": 1, "request": "blue headings", "target": "docx", "script": "edits/01.py", "status": "applied"},
            {"n": 2, "request": "tidy up", "target": "docx", "script": "edits/02.py", "status": "applied"},
        ),
    )
    job = Job("rebuild")
    notes = _emit(job, context, [_outcome()], RunUsage(model="x"))

    statuses = [e["status"] for e in context.presentation.edits]
    assert statuses == ["applied", "stale"]
    assert len(notes) == 1 and "tidy up" in notes[0] and "Ask again" in notes[0]
    report = job.artifacts["docx"]
    assert "Submitted on time." in _texts(report)
    # The zip is built after the replay, so it carries the edited report.
    with zipfile.ZipFile(job.artifacts["zip"]) as zf:
        inner = tmp_path / "from_zip.docx"
        inner.write_bytes(zf.read(report.name))
    assert "Submitted on time." in _texts(inner)
