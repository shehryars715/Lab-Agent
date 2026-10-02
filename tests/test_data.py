"""Datasets: profiling, resolution, placement, and the guards on fetching.

Offline, like the rest of the suite. Nothing here touches the network -- the
URL tests assert on what is REFUSED, which is decided before any socket opens,
and the Kaggle test asserts on the message you get without credentials.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from labsagent.data import (
    Acquisition,
    DATA_SUFFIXES,
    Dataset,
    describe,
    materialize,
    provenance_block,
)
from labsagent.data.preview import profile
from labsagent.data.sources import acquire, classify, from_path, kaggle_slug, resolve
from labsagent.errors import DataError

IRIS = (
    "sepal_length,sepal_width,species\n"
    "5.1,3.5,setosa\n"
    "4.9,3.0,setosa\n"
    "6.2,3.4,virginica\n"
    "5.9,3.0,virginica\n"
)


@pytest.fixture
def csv(tmp_path: Path) -> Path:
    path = tmp_path / "iris.csv"
    path.write_text(IRIS, encoding="utf-8")
    return path


# --- profiling --------------------------------------------------------------


def test_profile_names_every_column_and_its_type(csv: Path):
    """The whole point: the solver must never have to guess a column name."""
    text = profile(csv)
    assert "iris.csv" in text
    for column in ("sepal_length", "sepal_width", "species"):
        assert column in text
    assert "4 rows x 3 columns" in text
    # dtypes are what stop `df["sepal_length"].str.strip()` being written
    assert "float64" in text
    assert "setosa" in text, "sample rows reveal value formats, not just names"


def test_profile_of_a_non_table_still_describes_it(tmp_path: Path):
    """A file we cannot parse is still worth naming. It must not raise."""
    path = tmp_path / "notes.txt"
    path.write_text("this is prose, not a table", encoding="utf-8")
    text = profile(path)
    assert "notes.txt" in text
    assert "prose" in text


def test_profile_of_a_missing_file_does_not_raise(tmp_path: Path):
    assert "unreadable" in profile(tmp_path / "gone.csv")


# --- describing to the solver ----------------------------------------------


def test_describe_is_empty_without_data():
    """No data must cost the prompt nothing -- not even a header."""
    assert describe([]) == ""


def test_describe_forbids_the_two_expensive_mistakes(csv: Path):
    text = describe([Dataset(name="iris.csv", path=csv, preview=profile(csv))])
    assert "iris.csv" in text
    assert "read_file" in text, "a 64KB-truncated CSV read is the expensive failure"
    assert "download" in text


# --- placement --------------------------------------------------------------


def test_materialize_copies_into_the_workspace(tmp_path: Path, csv: Path):
    workspace = tmp_path / "workspace" / "task1"
    written = materialize([Dataset(name="iris.csv", path=csv)], workspace)
    assert written == ["iris.csv"]
    assert (workspace / "iris.csv").read_text(encoding="utf-8") == IRIS


def test_materialize_is_idempotent(tmp_path: Path, csv: Path):
    """Three attempts at one task must not recopy the file three times."""
    workspace = tmp_path / "ws"
    dataset = Dataset(name="iris.csv", path=csv)
    materialize([dataset], workspace)
    stamp = (workspace / "iris.csv").stat().st_mtime_ns
    materialize([dataset], workspace)
    assert (workspace / "iris.csv").stat().st_mtime_ns == stamp


def test_materialize_skips_a_file_that_vanished(tmp_path: Path):
    """A missing source is not a crashed run."""
    assert materialize([Dataset(name="x.csv", path=tmp_path / "nope.csv")], tmp_path) == []


# --- the report line --------------------------------------------------------


def test_provenance_names_the_source(csv: Path):
    block = provenance_block(
        [Dataset(name="iris.csv", path=csv, origin="kaggle", ref="uciml/iris")]
    )
    assert block is not None
    assert block.kind == "prose", "prose, so every emitter renders it in its own idiom"
    assert "iris.csv" in block.text
    assert "uciml/iris" in block.text


def test_provenance_is_none_without_data():
    assert provenance_block([]) is None


# --- classification ---------------------------------------------------------


def test_kaggle_slug_from_bare_slug_and_url():
    assert kaggle_slug("uciml/iris") == "uciml/iris"
    assert kaggle_slug("https://www.kaggle.com/datasets/uciml/iris") == "uciml/iris"
    assert kaggle_slug("https://www.kaggle.com/datasets/uciml/iris/") == "uciml/iris"


def test_a_competition_url_is_not_a_dataset_slug():
    """Competitions use a different API. Guessing would send a wrong request."""
    assert kaggle_slug("https://www.kaggle.com/c/titanic") is None


def test_an_existing_path_beats_the_kaggle_slug_shape(tmp_path: Path, monkeypatch):
    """`data/train.csv` looks exactly like `owner/name`. The filesystem decides."""
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "train.csv").write_text(IRIS, encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    assert classify("data/train.csv") == "path"


def test_an_unusable_reference_says_what_it_accepts():
    with pytest.raises(DataError) as exc:
        classify("just some words")
    assert "Kaggle" in str(exc.value)


# --- the URL guard ----------------------------------------------------------
#
# Every case below is refused BEFORE a socket is opened, so these run offline.


@pytest.mark.parametrize(
    "url",
    [
        "file:///C:/Users/HP/.env",
        "ftp://example.com/data.csv",
        "http://127.0.0.1:8000/data.csv",
        "http://localhost/data.csv",
        "http://169.254.169.254/latest/meta-data/",  # cloud metadata
        "http://10.0.0.5/data.csv",
        "http://192.168.1.1/data.csv",
    ],
)
def test_refuses_non_public_urls(url: str, tmp_path: Path):
    from labsagent.data.sources import from_url

    with pytest.raises(DataError):
        from_url(url, tmp_path)


def test_a_refused_url_is_a_failure_not_an_exception(tmp_path: Path):
    """`acquire` collects failures so one dead reference cannot sink the run."""
    result = acquire(["http://127.0.0.1/secret.csv"], tmp_path)
    assert result.datasets == []
    assert len(result.failures) == 1
    assert "127.0.0.1" in result.failures[0][1]


# --- local files ------------------------------------------------------------


def test_from_path_adopts_and_profiles(tmp_path: Path, csv: Path):
    dest = tmp_path / "data"
    found = from_path(str(csv), dest)
    assert len(found) == 1
    assert found[0].origin == "upload"
    assert (dest / "iris.csv").exists()
    assert "sepal_length" in found[0].preview


def test_a_python_file_is_not_data(tmp_path: Path):
    """`.py` is a lab format, not a dataset. Accepting it would let the solver
    be handed its own language as input."""
    script = tmp_path / "solution.py"
    script.write_text("print(1)", encoding="utf-8")
    assert ".py" not in DATA_SUFFIXES
    with pytest.raises(DataError):
        from_path(str(script), tmp_path / "out")


def test_oversized_file_is_refused_with_its_size(tmp_path: Path, csv: Path):
    with pytest.raises(DataError) as exc:
        from_path(str(csv), tmp_path / "out", max_bytes=1)
    assert "limit" in str(exc.value)


def test_resolve_dispatches_a_local_path(tmp_path: Path, csv: Path):
    found = resolve(str(csv), tmp_path / "out")
    assert [d.name for d in found] == ["iris.csv"]


# --- archives ---------------------------------------------------------------


def test_zip_yields_its_data_files(tmp_path: Path):
    archive = tmp_path / "bundle.zip"
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("train.csv", IRIS)
        z.writestr("readme.pdf", "ignored")
    found = from_path(str(archive), tmp_path / "out")
    assert [d.name for d in found] == ["train.csv"]


def test_zip_slip_is_refused(tmp_path: Path):
    """A member named ../escape.csv must never be written outside the run."""
    archive = tmp_path / "evil.zip"
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("../../escaped.csv", IRIS)
    with pytest.raises(DataError) as exc:
        from_path(str(archive), tmp_path / "out")
    assert "outside" in str(exc.value)
    assert not (tmp_path.parent / "escaped.csv").exists()


# --- Kaggle without credentials ---------------------------------------------


def test_kaggle_without_credentials_says_what_to_do(tmp_path: Path, monkeypatch):
    from labsagent.data import sources

    monkeypatch.delenv("KAGGLE_KEY", raising=False)
    monkeypatch.setattr(sources.Path, "home", staticmethod(lambda: tmp_path))

    with pytest.raises(DataError) as exc:
        sources.from_kaggle("uciml/iris", tmp_path / "out")
    message = str(exc.value)
    assert "KAGGLE_KEY" in message
    assert "attach" in message, "an unfixable error is not an actionable one"


# --- the prompt seam --------------------------------------------------------


def test_task_prompt_is_unchanged_without_data():
    """The cheap path must stay byte-identical: most labs have no data."""
    from labsagent.models import Task
    from labsagent.orchestrator import build_task_prompt

    task = Task(id="task1", title="Sum", statement="Read two ints and print the sum.")
    assert build_task_prompt(task, {}) == build_task_prompt(task, {}, [])


def test_task_prompt_carries_the_schema(csv: Path):
    from labsagent.models import Task
    from labsagent.orchestrator import build_task_prompt

    task = Task(id="task1", title="Cluster", statement="Cluster the data into 2 groups.")
    prompt = build_task_prompt(
        task, {}, [Dataset(name="iris.csv", path=csv, preview=profile(csv))]
    )
    assert "sepal_length" in prompt
    # The filename instruction is a structural invariant and stays LAST, after
    # the data block -- see the ordering rule in build_task_prompt.
    assert prompt.rindex("task1.py") > prompt.rindex("iris.csv")


# --- how the provenance line renders in each format -------------------------


def _outcome_with(note):
    from labsagent.blocks import blocks_for
    from labsagent.models import Task, TaskOutcome

    outcome = TaskOutcome(
        task=Task(id="task1", title="Cluster", statement="Cluster the data."),
        status="passed",
        code_text="import pandas as pd",
    )
    outcome.blocks = [note, *blocks_for(outcome)]
    return outcome


def test_an_upload_gets_no_parenthetical(csv: Path):
    """"supplied by you" reads wrong in a file the student hands to a grader."""
    note = provenance_block([Dataset(name="iris.csv", path=csv, origin="upload")])
    assert note.text == "Data used: iris.csv."


def test_the_py_carries_the_data_line_as_a_comment(csv: Path, tmp_path: Path):
    """PINNED DELIBERATELY. `emit/script.py` renders prose blocks as comments,
    so this line does reach the submitted script -- which is correct: a .py
    that calls read_csv is more useful for naming the file it needs. It is
    pinned because the wording is chosen to read well HERE as well as in the
    report, and a change to either side should have to say so."""
    from labsagent.emit import EmitContext
    from labsagent.emit.script import render
    from labsagent.models import LabSpec, Task

    note = provenance_block(
        [Dataset(name="iris.csv", path=csv, origin="kaggle", ref="uciml/iris")]
    )
    ctx = EmitContext(
        spec=LabSpec(lab_number="13", title="K-Means", tasks=[Task("task1", "C", "S")]),
        outcomes=[_outcome_with(note)],
        out_dir=tmp_path,
    )
    text = render(ctx)
    assert "# Data used: iris.csv (from Kaggle dataset uciml/iris)." in text
    # and it must still be a runnable script, not prose with code in it
    assert "import pandas as pd" in text


def test_the_markdown_carries_it_as_a_paragraph(csv: Path, tmp_path: Path):
    from labsagent.emit import EmitContext
    from labsagent.emit.markdown import render
    from labsagent.models import LabSpec, Task

    note = provenance_block([Dataset(name="iris.csv", path=csv, origin="upload")])
    ctx = EmitContext(
        spec=LabSpec(lab_number="13", title="K-Means", tasks=[Task("task1", "C", "S")]),
        outcomes=[_outcome_with(note)],
        out_dir=tmp_path,
    )
    text = render(ctx)
    assert "Data used: iris.csv." in text
    assert "# Data used" not in text, "a comment marker belongs only in the .py"


# --- collisions -------------------------------------------------------------


def test_two_references_producing_the_same_name_do_not_overwrite(tmp_path: Path):
    """Each resolve() starts with its own `taken` set, so uniqueness has to be
    checked against the DIRECTORY or the second file silently replaces the
    first -- and the run then solves against one file while naming two."""
    first = tmp_path / "a" / "data.csv"
    second = tmp_path / "b" / "data.csv"
    for path, body in ((first, IRIS), (second, IRIS + "6.0,3.1,virginica\n")):
        path.parent.mkdir(parents=True)
        path.write_text(body, encoding="utf-8")

    dest = tmp_path / "out"
    result = acquire([str(first), str(second)], dest)

    assert [d.name for d in result.datasets] == ["data.csv", "data_2.csv"]
    assert (dest / "data.csv").read_text(encoding="utf-8") == IRIS
    assert (dest / "data_2.csv").read_text(encoding="utf-8").endswith("6.0,3.1,virginica\n")


def test_a_malformed_url_is_a_dataerror_not_a_valueerror(tmp_path: Path):
    """urlsplit defers port parsing, so the ValueError surfaces at .port."""
    from labsagent.data.sources import from_url

    with pytest.raises(DataError):
        from_url("http://example.com:99999/data.csv", tmp_path)


def test_the_notebook_carries_it_as_a_markdown_cell(csv: Path, tmp_path: Path):
    """`package/notebook.py` reads named fields, not blocks, so it is the one
    emitter a blocks-only addition does NOT reach for free. A notebook that
    calls read_csv without naming the file is the least reproducible format."""
    import json

    from labsagent.emit import EmitContext
    from labsagent.emit.notebook import NotebookEmitter
    from labsagent.models import LabSpec, Task

    note = provenance_block(
        [Dataset(name="iris.csv", path=csv, origin="kaggle", ref="uciml/iris")]
    )
    ctx = EmitContext(
        spec=LabSpec(lab_number="13", title="K-Means", tasks=[Task("task1", "C", "S")]),
        outcomes=[_outcome_with(note)],
        out_dir=tmp_path,
    )
    (written,) = NotebookEmitter().emit(ctx)
    nb = json.loads(written.read_text(encoding="utf-8"))
    markdown = [
        "".join(c["source"]) for c in nb["cells"] if c["cell_type"] == "markdown"
    ]
    assert any("Data used: iris.csv (from Kaggle dataset uciml/iris)." == m for m in markdown)


def test_the_notebook_does_not_duplicate_the_explanation(csv: Path, tmp_path: Path):
    """`blocks_for()` synthesises prose LAST, so a naive block walk here would
    print every explanation twice. Only leading prose is taken."""
    import json

    from labsagent.emit import EmitContext
    from labsagent.emit.notebook import NotebookEmitter
    from labsagent.models import LabSpec, Task, TaskOutcome
    from labsagent.present import AnswerPlan, Include

    outcome = TaskOutcome(
        task=Task(id="task1", title="C", statement="S"),
        status="passed",
        code_text="print(1)",
        explanation="It prints one.",
    )
    ctx = EmitContext(
        spec=LabSpec(lab_number="1", title="L", tasks=[outcome.task]),
        outcomes=[outcome],
        out_dir=tmp_path,
        # Explanations are left out unless asked for (2026-10-02); this test is
        # about one that IS asked for not being printed twice.
        plans={"task1": AnswerPlan(include=Include(explanation=True))},
    )
    (written,) = NotebookEmitter().emit(ctx)
    source = json.dumps(json.loads(written.read_text(encoding="utf-8")))
    assert source.count("It prints one.") == 1


def test_the_annotated_docx_carries_it_above_the_code(csv: Path, tmp_path: Path):
    """The annotate-in-place report is the flagship deliverable for a .docx
    upload, and it reads TaskOutcome's named fields rather than blocks -- so it
    was the third emitter a blocks-only addition silently missed."""
    import docx as pydocx

    from labsagent.models import LabSpec, Task
    from labsagent.report.docx_builder import annotate_manual

    source = tmp_path / "manual.docx"
    doc = pydocx.Document()
    doc.add_paragraph("Task 1: Cluster the data.")
    doc.save(str(source))

    note = provenance_block(
        [Dataset(name="iris.csv", path=csv, origin="kaggle", ref="uciml/iris")]
    )
    out = annotate_manual(
        source,
        tmp_path / "report.docx",
        [_outcome_with(note)],
        anchors={"task1": 0},
    )
    text = [p.text for p in pydocx.Document(str(out)).paragraphs]
    assert "Data used: iris.csv (from Kaggle dataset uciml/iris)." in text
    assert text.index("Data used: iris.csv (from Kaggle dataset uciml/iris).") < text.index(
        "Code:"
    ), "provenance belongs above the code, not buried after it"


def test_the_annotated_docx_is_unchanged_without_blocks(tmp_path: Path):
    """Byte-stability of the anchored path is what emit/docx.py promises."""
    import docx as pydocx

    from labsagent.models import Task, TaskOutcome
    from labsagent.report.docx_builder import annotate_manual

    def render(target: str) -> list[str]:
        source = tmp_path / f"manual_{target}.docx"
        doc = pydocx.Document()
        doc.add_paragraph("Task 1: Do the thing.")
        doc.save(str(source))
        outcome = TaskOutcome(
            task=Task(id="task1", title="T", statement="S"),
            status="passed",
            code_text="print(1)",
        )
        out = annotate_manual(
            source, tmp_path / f"{target}.docx", [outcome], anchors={"task1": 0}
        )
        return [p.text for p in pydocx.Document(str(out)).paragraphs]

    assert render("a") == render("b")
    assert not any("Data used" in line for line in render("a"))


# --- when nothing resolved --------------------------------------------------
#
# The system prompt tells the solver that listed data files "are already saved
# in your workspace". With no data and no block, it read that, believed it, and
# went looking outside its workspace for a file it had never been given.


def test_describe_says_so_when_every_reference_failed():
    text = describe([], [("Online", "I do not know how to get 'Online'.")])
    assert "No data was resolved" in text
    assert "Online" in text
    assert "already saved in your workspace" not in text
    assert "do not invent substitute data" in text.lower()


def test_describe_is_still_empty_for_a_lab_that_wanted_no_data():
    """A lab with no data must stay byte-identical to what it was."""
    assert describe([]) == ""
    assert describe([], []) == ""


def test_a_lab_that_asked_for_nothing_is_not_a_total_failure():
    assert Acquisition().total_failure is False


def test_asking_and_getting_nothing_is_a_total_failure():
    acquired = Acquisition(requested=["a.csv"], failures=[("a.csv", "gone")])
    assert acquired.total_failure is True


def test_one_failure_among_several_is_not_a_total_failure(csv: Path):
    acquired = Acquisition(
        datasets=[Dataset(name="a.csv", path=csv)],
        requested=["a.csv", "b.csv"],
        failures=[("b.csv", "gone")],
    )
    assert acquired.total_failure is False, "the partial case must keep running"


def test_acquire_records_what_it_was_asked_for(tmp_path: Path, csv: Path):
    got = acquire([str(csv), str(csv)], tmp_path / "data")
    assert got.requested == [str(csv)], "deduped, and recorded even when it worked"


# --- workbooks --------------------------------------------------------------
#
# ".xlsx" was an accepted data suffix that nothing could read: openpyxl was not
# a declared dependency, so `pd.read_excel` raised, the profile fell back to
# reading the zip container as text, and the solver -- told the file was in its
# workspace -- spent most of a run writing its own xlsx parser.


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "sales.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Q1"
    sheet.append(["region", "units"])
    sheet.append(["north", 4])
    sheet.append(["south", 7])
    book.save(path)
    return path


def test_a_workbook_is_profiled_by_its_columns_not_its_bytes(workbook: Path):
    text = profile(workbook)
    assert "region" in text and "units" in text
    assert "PK" not in text, "the zip container must never reach a prompt"


def test_a_binary_file_is_refused_rather_than_rendered(tmp_path: Path):
    path = tmp_path / "mystery.bin"
    path.write_bytes(b"PK\x03\x04" + b"\xff\xfe" * 200)
    text = profile(path)
    assert "binary file" in text
    assert "�" not in text, "no replacement characters, which is what this cost"


def test_a_workbook_arrives_as_a_csv(tmp_path: Path, workbook: Path):
    """The solver never meets the workbook: it is converted once, here."""
    datasets = from_path(str(workbook), tmp_path / "data")
    assert [d.name for d in datasets] == ["sales.csv"]
    assert datasets[0].path.suffix == ".csv"
    assert datasets[0].path.read_text(encoding="utf-8").startswith("region,units")
    assert "converted from sales.xlsx" in datasets[0].preview
    assert datasets[0].ref == "sales.xlsx", "provenance survives the conversion"


def test_the_original_workbook_is_kept_but_not_handed_over(tmp_path: Path, workbook: Path):
    data_dir = tmp_path / "data"
    datasets = from_path(str(workbook), data_dir)
    assert (data_dir / "sales.xlsx").exists(), "the run keeps a record of its input"
    assert "sales.xlsx" not in [d.name for d in datasets]


def test_conversion_can_be_turned_off(tmp_path: Path, workbook: Path):
    datasets = from_path(str(workbook), tmp_path / "data", convert_excel=False)
    assert [d.name for d in datasets] == ["sales.xlsx"]


def test_a_multi_sheet_workbook_names_each_sheet(tmp_path: Path):
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "book.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "First"
    book.active.append(["a"])
    second = book.create_sheet("Second Half")
    second.append(["b"])
    book.save(path)

    names = [d.name for d in from_path(str(path), tmp_path / "data")]
    assert names == ["book__first.csv", "book__second_half.csv"]


def test_an_unreadable_workbook_falls_back_to_the_file_itself(tmp_path: Path):
    """Conversion is an optimisation; it must never cost the dataset."""
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"PK\x03\x04not really a workbook")
    datasets = from_path(str(path), tmp_path / "data")
    assert [d.name for d in datasets] == ["broken.xlsx"]


def test_to_csv_on_rubbish_returns_nothing(tmp_path: Path):
    from labsagent.data.normalise import to_csv

    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a zip at all")
    assert to_csv(path, tmp_path / "out") == []
