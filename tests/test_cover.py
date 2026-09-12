"""The cover sheet: prepended without disturbing the manual it fronts."""

from __future__ import annotations

import docx
import pytest
from docx.enum.text import WD_BREAK

from labsagent.models import LabSpec, Task
from labsagent.profile import StudentProfile
from labsagent.ingest.cover import CoverFacts
from labsagent.report.cover import CoverInfo, build_cover, cover_from


@pytest.fixture
def manual(tmp_path):
    doc = docx.Document()
    doc.add_heading("Original Lab Manual", level=1)
    doc.add_paragraph("Task 1: do the thing")
    doc.add_paragraph("Some closing text")
    path = tmp_path / "manual.docx"
    doc.save(str(path))
    return path


FULL = CoverInfo(
    lab_number="10",
    lab_title="Neural Networks-Part A",
    course="CS245 Machine Learning",
    student_name="Shehryar",
    cms_id="22F-1234",
    section="BSDS-02A",
    instructor="Dr. Nazia Perwaiz",
    lab_engineer="Alishba Zulfiqar",
    date="8th April 2026",
    department="Faculty of Computing",
)


class TestRows:
    def test_includes_the_four_the_user_asked_for(self):
        labels = dict(FULL.rows())
        assert labels["Name"] == "Shehryar"
        assert labels["CMS ID"] == "22F-1234"
        assert labels["Lab Number"] == "Lab 10"
        assert labels["Instructor"] == "Dr. Nazia Perwaiz"

    def test_unknown_fields_are_omitted_not_blank(self):
        """A row reading 'Instructor:' with nothing after it looks broken."""
        labels = dict(CoverInfo(student_name="S", cms_id="1").rows())
        assert "Instructor" not in labels
        assert "Date" not in labels

    def test_whitespace_only_counts_as_unknown(self):
        assert "Instructor" not in dict(CoverInfo(instructor="   ").rows())

    def test_name_comes_first(self):
        assert FULL.rows()[0][0] == "Name"


class TestBuildCover:
    def test_cover_lands_before_the_manual(self, manual):
        doc = docx.Document(str(manual))
        build_cover(doc, FULL)
        texts = [p.text for p in doc.paragraphs]
        assert texts.index("Lab 10: Neural Networks-Part A") < texts.index(
            "Original Lab Manual"
        )

    def test_original_content_survives_verbatim(self, manual):
        before = [p.text for p in docx.Document(str(manual)).paragraphs if p.text.strip()]
        doc = docx.Document(str(manual))
        build_cover(doc, FULL)
        after = [p.text for p in doc.paragraphs]
        assert all(text in after for text in before)

    def test_fields_land_in_a_table(self, manual):
        doc = docx.Document(str(manual))
        build_cover(doc, FULL)
        assert doc.tables, "cover fields should be tabulated"
        cells = {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
        assert cells["CMS ID"] == "22F-1234"
        assert cells["Lab Engineer"] == "Alishba Zulfiqar"

    def test_a_page_break_separates_cover_from_manual(self, manual):
        doc = docx.Document(str(manual))
        build_cover(doc, FULL)
        breaks = [
            br
            for p in doc.paragraphs
            for r in p.runs
            for br in r._element.findall(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br"
            )
        ]
        assert breaks, "expected a page break after the cover"

    def test_survives_a_manual_with_no_known_fields(self, manual):
        doc = docx.Document(str(manual))
        build_cover(doc, CoverInfo())
        assert "Lab Report" in [p.text for p in doc.paragraphs]

    def test_cover_order_is_preserved_not_reversed(self, manual):
        """The same reverse-insert trap the annotator hits, on a different path."""
        doc = docx.Document(str(manual))
        build_cover(doc, FULL)
        texts = [p.text for p in doc.paragraphs]
        assert texts.index("FACULTY OF COMPUTING") < texts.index(
            "Lab 10: Neural Networks-Part A"
        )
        assert texts.index("Lab 10: Neural Networks-Part A") < texts.index(
            "CS245 Machine Learning"
        )


class TestCoverFrom:
    def test_merges_spec_profile_and_manual(self):
        spec = LabSpec(lab_number="13", title="K-Means", course=None, tasks=[Task(id="t", title="t", statement="s")])
        facts = CoverFacts(course="CS245 ML", instructor="Dr. N", date="1 Jan 2026")
        info = cover_from(spec, StudentProfile(name="S", cms_id="22F-1"), facts)
        assert info.lab_number == "13"
        assert info.course == "CS245 ML"
        assert info.instructor == "Dr. N"
        assert info.student_name == "S"

    def test_profile_section_beats_the_manual(self):
        spec = LabSpec(lab_number="1", title="T", tasks=[])
        facts = CoverFacts(section="OLD-01")
        info = cover_from(spec, StudentProfile(name="S", cms_id="1", section="NEW-02"), facts)
        assert info.section == "NEW-02"

    def test_manual_section_used_when_profile_is_silent(self):
        spec = LabSpec(lab_number="1", title="T", tasks=[])
        info = cover_from(spec, StudentProfile(name="S", cms_id="1"), CoverFacts(section="BSDS-02A"))
        assert info.section == "BSDS-02A"
